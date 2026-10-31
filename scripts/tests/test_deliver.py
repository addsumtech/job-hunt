"""`deliver.py` — the hand-off step, and the ordering it belongs to.

A workspace under `~/.claude/job-profiles/` is where the skill works and not
where a person looks. These tests pin what gets handed over, what deliberately
does not, and the one thing a PDF pipeline must never do quietly: drop the
characters it could not render.
"""
import json
import pathlib
import shutil
import subprocess
import sys

import pytest

REPO = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "scripts"))

import deliver  # noqa: E402

HAVE_PDF = bool(shutil.which("pandoc") and shutil.which("tectonic")
                and shutil.which("pdftotext"))


def build(tmp_path, md="# rows\n"):
    ws = tmp_path / "2026-09-06-round"
    (ws / "raw" / "opencli-help").mkdir(parents=True)
    (ws / "mock").mkdir()
    (ws / "report.md").write_text(md, encoding="utf-8")
    (ws / "shortlist.yaml").write_text("rows: []\n", encoding="utf-8")
    import pymupdf
    with pymupdf.open() as doc:
        doc.new_page().insert_text((72, 72), "Candidate CV")
        doc.save(ws / "cv.pdf")
    (ws / "mock" / "assessment-1.md").write_text("# round 1\n", encoding="utf-8")
    (ws / "journal.jsonl").write_text('{"action":"gate"}\n', encoding="utf-8")
    (ws / "raw" / "51job-1.json").write_text("[]", encoding="utf-8")
    (ws / "51job-1.err").write_text("boom\n", encoding="utf-8")
    (ws / ".DS_Store").write_bytes(b"\x00")
    return ws


def run(ws, dest, *extra):
    return deliver.main(["--workspace", str(ws), "--to", str(dest), *extra])


# Spelled out as well as imported, so these tests still run (and fail for the
# real reason) against a deliver.py that predates the manifest.
MANIFEST = getattr(deliver, "MANIFEST_NAME", ".jobhunt-delivery.json")


def client_files(dest):
    """What the client sees: every file except the hidden delivery manifest."""
    return {p.relative_to(dest).as_posix() for p in dest.rglob("*")
            if p.is_file() and p.name != MANIFEST}


# ---- what lands, and where ------------------------------------------------

def test_client_documents_share_one_explicit_folder(tmp_path):
    ws = build(tmp_path)
    dest = tmp_path / "consultation"
    assert run(ws, dest, "--no-pdf") == 0
    assert {p.name for p in dest.iterdir()} == {"简历", "报告", MANIFEST}
    assert client_files(dest) == {"报告/求职建议报告.md", "简历/简历.pdf"}


def test_default_destination_is_a_consultation_folder(tmp_path, monkeypatch):
    ws = build(tmp_path)
    monkeypatch.setattr(deliver, "DEFAULT_ROOT", tmp_path / "Downloads")
    assert deliver.main(["--workspace", str(ws), "--no-pdf"]) == 0
    assert (tmp_path / "Downloads" / ws.name / "报告" / "求职建议报告.md").is_file()


def test_apply_delivers_its_readiness_brief_without_private_review_records(tmp_path):
    ws = build(tmp_path)
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "apply"})
    brief = "# 面试准备\n\n说明你本人负责的需求整理和验收工作。\n"
    (ws / "interview-brief.md").write_text(brief, encoding="utf-8")
    (ws / "claims.yaml").write_text("private provenance records")
    dest = tmp_path / "consultation"
    assert run(ws, dest, "--no-pdf") == 0
    assert (dest / "面试准备/面试准备.md").read_text(encoding="utf-8") == brief
    assert not list(dest.rglob("claims*"))
    assert not list(dest.rglob("assessment*"))


@pytest.mark.parametrize("mode", ["assess", "interview"])
def test_read_only_consultations_do_not_redeliver_or_render_the_source_cv(tmp_path, mode):
    ws = build(tmp_path, "# Career report\n\nDiscuss the documented experience.\n")
    (ws / "cv.md").write_text("# Source CV\n\nAn existing source document.", encoding="utf-8")
    (ws / "interview-brief.md").write_text("# 面试准备\n\n说明你本人负责的部分。\n",
                                           encoding="utf-8")
    deliver.journal.append(ws, {"action": "mode_entry", "mode": mode})
    dest = tmp_path / "consultation"
    assert run(ws, dest) == 0
    brief = {"面试准备/面试准备.md", "面试准备/面试准备.pdf"} if mode == "interview" else set()
    assert client_files(dest) == {"报告/求职建议报告.md", "报告/求职建议报告.pdf"} | brief
    assert not (dest / "简历").exists()
    assert (ws / "cv.md").is_file() and (ws / "cv.pdf").is_file()


def test_an_interview_redelivers_the_brief_with_its_walk_back_list(tmp_path):
    """The walk-backs are agreed in the interview round; the copy apply delivered
    does not have them, so the debrief folder must carry the current brief."""
    ws = build(tmp_path)
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "apply"})
    (ws / "interview-brief.md").write_text("# 面试准备\n\n说明你本人负责的部分。\n",
                                           encoding="utf-8")
    dest = tmp_path / "consultation"
    assert run(ws, dest, "--no-pdf") == 0
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "interview"})
    updated = "# 面试准备\n\n说明你本人负责的部分。\n\n## Walk-back list\n\n- 数据室不是我全权负责。\n"
    (ws / "interview-brief.md").write_text(updated, encoding="utf-8")
    assert run(ws, dest, "--no-pdf") == 0
    delivered = [p for p in dest.rglob("面试准备.md")
                 if p.read_text(encoding="utf-8") == updated]
    assert delivered, sorted(str(p.relative_to(dest)) for p in dest.rglob("*.md"))


def test_explicit_application_copy_remains_available_after_an_interview(tmp_path):
    ws = build(tmp_path)
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "interview"})
    dest = tmp_path / "requested-copy"
    assert run(ws, dest, "--no-pdf", "--include-applications") == 0
    assert (dest / "简历" / "简历.pdf").read_bytes() == (ws / "cv.pdf").read_bytes()


# ---- what stays behind ----------------------------------------------------

def test_the_provenance_chain_is_left_in_the_workspace(tmp_path):
    """`raw/` and the receipts are what an audit reads IN PLACE.

    Copying them makes a second, drifting copy of the only evidence any claim in
    these documents is traced to — and it is the bulk of the bytes.
    """
    ws = build(tmp_path)
    dest = tmp_path / "out"
    assert run(ws, dest, "--no-pdf") == 0
    names = {p.name for p in dest.iterdir()}
    assert not any("journal" in n or n.endswith(".err") or "51job-1.json" in n
                   for n in names), names
    assert (ws / "raw" / "51job-1.json").is_file()
    assert (ws / "journal.jsonl").is_file()


def test_it_is_a_copy_and_never_a_move(tmp_path):
    """`paths.py` owns that layout, apply mode's resume lookup finds work BY the
    path shape, and `check_claims.py` fingerprints the master profile there."""
    ws = build(tmp_path)
    before = sorted(p.relative_to(ws) for p in ws.rglob("*") if p.is_file())
    assert run(ws, tmp_path / "out", "--no-pdf") == 0
    assert before == sorted(p.relative_to(ws) for p in ws.rglob("*") if p.is_file())


# ---- the PDF, which is verified rather than trusted -----------------------

def test_a_pdf_that_dropped_characters_is_deleted_not_delivered(tmp_path, monkeypatch):
    """MEASURED: `pandoc --pdf-engine=tectonic` on a Chinese document exits 0 and
    writes a PDF whose every CJK glyph is a box — 528 characters in, 0 read back.

    A wrong artifact behind a green exit code is the failure this whole check
    exists for, so the round-trip is the gate and the exit code is not.
    """
    ws = build(tmp_path, md="# 岗位候选\n\n这是中文内容。\n")
    dest = tmp_path / "out"
    monkeypatch.setattr(deliver, "_render_portable_report",
                        lambda *args: (False, "disabled for this regression"))
    monkeypatch.setattr(deliver, "pick_cjk_font", lambda *args: "SomeFont")
    monkeypatch.setattr(deliver, "_pandoc",
                        lambda md, pdf, font: (pdf.write_bytes(b"%PDF"), True)[1])
    monkeypatch.setattr(deliver, "pdf_text", lambda pdf: "boxes only, no CJK")
    assert run(ws, dest) == 2
    assert not (dest / "报告" / "求职建议报告.pdf").exists()
    assert (dest / "报告" / "求职建议报告.md").is_file(), "the Markdown still ships"


def test_pdf_symbol_conversion_changes_only_pandoc_text_nodes():
    document = {
        "blocks": [{"t": "Para", "c": [
            {"t": "Str", "c": "\u26a0\ufe0f Warning"},
            {"t": "Code", "c": [["", [], []], "\u26a0 keep code exact"]},
            {"t": "Link", "c": [
                ["", [], []], [{"t": "Str", "c": "\u26a0 link"}],
                ["https://example.test/\u26a0", ""],
            ]},
        ]}],
    }
    assert deliver._replace_unsupported_pdf_symbols(document)
    para = document["blocks"][0]["c"]
    assert para[0]["c"] == "[!] Warning"
    assert para[1]["c"][1] == "\u26a0 keep code exact"
    assert para[2]["c"][1][0]["c"] == "[!] link"
    assert para[2]["c"][2][0] == "https://example.test/\u26a0"


def test_pandoc_symbol_ast_keeps_a_warning_used_only_as_code(tmp_path, monkeypatch):
    md = tmp_path / "code.md"
    md.write_text("`\u26a0 keep code exact`\n", encoding="utf-8")
    document = {"blocks": [{"t": "Para", "c": [
        {"t": "Code", "c": [["", [], []], "\u26a0 keep code exact"]},
    ]}]}

    def pandoc_json(cmd, **kwargs):
        return subprocess.CompletedProcess(cmd, 0, json.dumps(document).encode(), b"")

    monkeypatch.setattr(deliver.subprocess, "run", pandoc_json)
    ast = deliver._pandoc_json_with_readable_symbols(md)
    assert ast is not None
    assert json.loads(ast)["blocks"][0]["c"][0]["c"][1] == "\u26a0 keep code exact"


def test_chinese_report_uses_bundled_cjk_renderer_when_pandoc_has_no_cjk_font(tmp_path, monkeypatch):
    ws = build(tmp_path, md=(
        "# 岗位候选\n\n这是中文内容。\n\n"
        "| 优先级 | 岗位 | 公司与地点 | 展示薪资 | 初判与投入 | 依据 |\n"
        "| --- | --- | --- | --- | --- | --- |\n"
        "| 1 | 大模型工程师 | 示例公司，上海 | 25-40K | worth_applying | 多模态模型经验 |\n"
    ))
    dest = tmp_path / "out"
    monkeypatch.setattr(deliver, "pick_cjk_font", lambda *args: None)

    real_run = deliver.subprocess.run

    def poppler_without_cjk_text(args, *positional, **keywords):
        if args and args[0] == "pdftotext":
            return subprocess.CompletedProcess(args, 0, b"", b"")
        return real_run(args, *positional, **keywords)

    monkeypatch.setattr(deliver.subprocess, "run", poppler_without_cjk_text)
    assert run(ws, dest) == 0
    pdf = dest / "报告" / "求职建议报告.pdf"
    assert pdf.is_file()
    assert "大模型工程师" in deliver.pdf_text(pdf)


def test_cjk_detection_and_counting():
    assert deliver.has_cjk("岗位候选")
    assert deliver.has_cjk("ソフトウェア")
    assert deliver.has_cjk("소프트웨어")
    assert not deliver.has_cjk("Senior ATD Engineer - System (Open)")
    assert deliver.cjk_chars("岗位候选 x3") == 4


def test_raw_source_urls_become_short_clickable_links_in_chinese_pdfs():
    rendered = deliver._inline_pdf_markup(
        "来源：<https://www.zhipin.com/job_detail/abc.html?ka=search_list&foo=bar>、"
        "<https://jobs.51job.com/shanghai/123456.html>"
    )
    assert ">zhipin.com</link>" in rendered
    assert "href=\"https://www.zhipin.com/job_detail/abc.html?ka=search_list&amp;foo=bar\"" in rendered
    assert ">jobs.51job.com</link>" in rendered
    assert "href=\"https://jobs.51job.com/shanghai/123456.html\"" in rendered
    assert "&lt;https" not in rendered


def test_chinese_pdf_contains_distinct_clickable_urls(tmp_path):
    first = "https://www.zhipin.com/job_detail/abc.html?ka=search_list&foo=bar"
    second = "https://jobs.51job.com/shanghai/123456.html"
    ws = build(tmp_path, md=f"# 岗位候选\n\n- <{first}>、<{second}>\n")
    dest = tmp_path / "out"
    assert run(ws, dest) == 0
    import pymupdf
    with pymupdf.open(dest / "报告" / "求职建议报告.pdf") as pdf:
        urls = {link.get("uri") for page in pdf for link in page.get_links()}
    assert {first, second} <= urls


def test_pdf_text_falls_back_to_pymupdf_for_bundled_cjk_fonts(tmp_path, monkeypatch):
    """A Poppler build that returns no text must not reject a readable PDF."""
    import pymupdf
    pdf = tmp_path / "cjk.pdf"
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "大模型工程师", fontname="china-s", fontsize=12)
        document.save(pdf)

    def blank_poppler(*args, **kwargs):
        # A failed subprocess commonly has stdout=None, not b"".
        return subprocess.CompletedProcess(args[0], 1)

    monkeypatch.setattr(deliver.subprocess, "run", blank_poppler)
    assert "大模型工程师" in deliver.pdf_text(pdf)


def test_pdf_with_an_authored_link_but_no_link_annotation_is_refused(tmp_path):
    """Blue-looking text is not an actionable link unless the PDF annotates it."""
    import pymupdf
    pdf = tmp_path / "unlinked.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "Open posting")
        document.save(pdf)

    ok, why = deliver._verify_pdf(
        pdf, "[Open posting](https://jobs.example.test/roles/123)", False)
    assert not ok
    assert "PDF_LINKS_MISSING" in why
    assert not pdf.exists(), "an unclickable report PDF must not be delivered"


@pytest.mark.parametrize("source", [
    "研究エンジニアと機械学習",
    "AI 연구개발 엔지니어",
    "Senior AI engineer with reliable deployment experience",
    "Ingeniería de aprendizaje automático y visión por computadora",
])
def test_localized_reports_use_the_bundled_portable_renderer(tmp_path, source):
    md, pdf = tmp_path / "report.md", tmp_path / "report.pdf"
    md.write_text(source, encoding="utf-8")
    assert deliver.render_pdf(md, pdf, None) == (True, "")
    assert source in deliver.pdf_text(pdf)


@pytest.mark.skipif(not HAVE_PDF, reason="needs pandoc + tectonic + pdftotext")
def test_a_chinese_document_really_round_trips_through_a_real_pdf(tmp_path):
    """The end-to-end one. Everything above can pass on a mocked renderer."""
    ws = build(tmp_path, md="# 岗位候选\n\n医学影像与图像重建方向的岗位清单。\n")
    dest = tmp_path / "out"
    assert run(ws, dest) == 0
    pdf = dest / "报告" / "求职建议报告.pdf"
    assert pdf.is_file()
    assert deliver.cjk_chars(deliver.pdf_text(pdf)) >= 15


# ---- could-not-run, never a silent success --------------------------------

def test_an_unwritable_destination_exits_2_and_copies_nothing(tmp_path, monkeypatch):
    """macOS TCC can start refusing ~/Downloads part-way through a session, and
    `os.access` says yes while the write fails — so the probe writes a real file."""
    ws = build(tmp_path)
    monkeypatch.setattr(deliver, "writable", lambda d: (False, "Operation not permitted"))
    dest = tmp_path / "blocked"
    assert deliver.main(["--workspace", str(ws), "--to", str(dest)]) == 2
    assert not dest.exists()


def test_a_missing_workspace_exits_2(tmp_path):
    assert deliver.main(["--workspace", str(tmp_path / "nope"),
                         "--to", str(tmp_path / "out")]) == 2


def test_delivering_into_the_workspace_is_refused(tmp_path):
    ws = build(tmp_path)
    assert deliver.main(["--workspace", str(ws), "--to", str(ws / "export")]) == 2
    assert deliver.main(["--workspace", str(ws), "--to", str(ws)]) == 2


def test_a_workspace_with_only_provenance_exits_2(tmp_path):
    ws = tmp_path / "empty-round"
    (ws / "raw").mkdir(parents=True)
    (ws / "raw" / "x.json").write_text("[]", encoding="utf-8")
    (ws / "journal.jsonl").write_text("{}\n", encoding="utf-8")
    assert deliver.main(["--workspace", str(ws), "--to", str(tmp_path / "out")]) == 2


def test_discover_delivery_requires_a_checked_shortlist(tmp_path):
    ws = build(tmp_path)
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "discover"})
    dest = tmp_path / "out"
    assert run(ws, dest, "--no-pdf") == 2
    assert not dest.exists()

    (ws / "shortlist.md").write_text("# Shortlist\n", encoding="utf-8")
    deliver.journal.receipt(ws, "check_shortlist", {}, "pass")
    # A checked empty round alone is not a completed search consultation.
    assert run(ws, dest, "--no-pdf") == 2
    (ws / "brief.yaml").write_text(
        "report_scope: preliminary\npreliminary_request: Show this initial round now\n")
    assert run(ws, dest, "--no-pdf") == 0


def test_discover_delivery_requires_every_shortlist_link_in_the_report(tmp_path):
    ws = build(tmp_path, md="# Client report\n")
    url = "https://jobs.example.test/roles/123?tracking=source"
    (ws / "shortlist.yaml").write_text(
        f"rows:\n  - url: {url}\n", encoding="utf-8")
    (ws / "shortlist.md").write_text("# Shortlist\n", encoding="utf-8")
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "discover"})
    deliver.journal.receipt(ws, "check_shortlist", {}, "pass")
    (ws / "brief.yaml").write_text(
        "report_scope: preliminary\npreliminary_request: Show me the initial shortlist first\n",
        encoding="utf-8")

    dest = tmp_path / "out"
    assert run(ws, dest, "--no-pdf") == 2
    assert not dest.exists()

    (ws / "report.md").write_text(
        f"# Client report\n\n[Open posting]({url})\n", encoding="utf-8")
    assert run(ws, dest, "--no-pdf") == 0


def test_discover_delivery_refuses_an_unreadable_shortlist(tmp_path):
    """Delivery must return its normal incomplete status, never parse-crash."""
    ws = build(tmp_path)
    (ws / "shortlist.yaml").write_text("rows: [unterminated", encoding="utf-8")
    (ws / "shortlist.md").write_text("# Shortlist\n", encoding="utf-8")
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "discover"})
    deliver.journal.receipt(ws, "check_shortlist", {}, "pass")

    dest = tmp_path / "out"
    assert run(ws, dest, "--no-pdf") == 2
    assert not dest.exists()


def test_it_runs_as_a_script(tmp_path):
    """The mode files tell the model to run a COMMAND. It has to work as one."""
    ws = build(tmp_path)
    dest = tmp_path / "out"
    r = subprocess.run([sys.executable, str(REPO / "scripts" / "deliver.py"),
                        "--workspace", str(ws), "--to", str(dest), "--no-pdf"],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stderr
    assert str(dest) in r.stdout, "the path to quote must be printed"


# ---- the skill has to tell the run to do this -----------------------------

@pytest.mark.parametrize("mode", ["discover", "assess", "apply", "interview"])
def test_every_mode_file_names_the_delivery_step(mode):
    """Prose in SKILL.md that only one mode repeats is prose three modes skip."""
    text = (REPO / "modes" / f"{mode}.md").read_text(encoding="utf-8")
    assert "deliver.py" in text, f"{mode}.md never tells the run to hand the files over"
    assert "~/Downloads" in text


def test_the_skill_file_lists_delivery_and_the_path_to_tell_the_user():
    text = (REPO / "SKILL.md").read_text(encoding="utf-8")
    assert "scripts/deliver.py" in text
    assert "~/Downloads" in text


# ---- flattening must not silently drop a file -----------------------------
#
# MEASURED on the first version of this script, by probing it rather than by
# reading it. `mock/notes.md` and a root-level `notes.md` both became
# `<slug>-notes.md`: one overwrote the other, and the run printed "Delivered 2
# file(s)" over a directory holding one. Interview mode writes
# `mock/transcript-n.md`, `mock/assessment-n.md` and `mock/answer-guide.md`
# beside root-level outputs, so this is the ordinary layout, not a corner.

def test_internal_reports_are_excluded_even_if_markdown(tmp_path):
    ws = build(tmp_path)
    (ws / "completion.md").write_text("TOOL BUGS")
    (ws / "mock" / "report.md").write_text("INTERNAL REVIEW")
    dest = tmp_path / "out"
    assert run(ws, dest, "--no-pdf") == 0
    assert not any("TOOL BUGS" in p.read_text() or "INTERNAL REVIEW" in p.read_text()
                   for p in dest.rglob("*.md"))


def test_the_reported_count_matches_what_actually_landed(tmp_path, capsys):
    """The overwrite was survivable; reporting two deliveries over one file was
    the part that would have gone unnoticed."""
    ws = tmp_path / "round"
    (ws / "mock").mkdir(parents=True)
    (ws / "report.md").write_text("ROOT\n", encoding="utf-8")
    (ws / "mock" / "notes.md").write_text("NESTED\n", encoding="utf-8")
    assert run(ws, tmp_path / "out", "--no-pdf") == 0
    said = capsys.readouterr().out
    landed = len(client_files(tmp_path / "out"))
    assert f"Delivered {landed} file(s)" in said


def test_flat_name_keeps_the_whole_relative_path():
    assert deliver.flat_name("r", pathlib.Path("mock/assessment-1.md")) == \
        "r-mock__assessment-1.md"
    assert deliver.flat_name("r", pathlib.Path("cv.md")) == "r-cv.md"


# ---- a hand-off that left no trace ----------------------------------------
#
# The gates run BEFORE delivery, so no composer can check that delivery
# happened — and a gate that says what to do next is prompting rather than
# checking, which this repo has decided against on purpose
# (test_mode_declaration_and_handoff.py: "a clean gate run still says nothing on
# stdout"). What is left is a record: a later reader can tell a run that skipped
# the last step from one that took it.

def test_a_delivery_leaves_a_record_in_the_workspace_journal(tmp_path):
    import json
    ws = build(tmp_path)
    assert run(ws, tmp_path / "out", "--no-pdf") == 0
    records = [json.loads(l) for l in
               (ws / "journal.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
    delivery = [r for r in records if r.get("action") == "delivery"]
    assert len(delivery) == 1
    assert delivery[0]["destination"].endswith("out")
    assert "报告/求职建议报告.md" in delivery[0]["files"]


def test_the_record_is_not_a_gate_receipt(tmp_path):
    """Delivery decides nothing and has no verdict. Filing it as a gate would
    put a passing receipt in the journal for something that was never a check —
    and `check_apply` fails on ANY gate receipt reading `fail`, named or not."""
    import json

    def gates(ws):
        return [json.loads(l) for l
                in (ws / "journal.jsonl").read_text(encoding="utf-8").splitlines()
                if l.strip() and json.loads(l).get("action") == "gate"]

    ws = build(tmp_path)              # the fixture seeds one unrelated gate line
    before = len(gates(ws))
    assert run(ws, tmp_path / "out", "--no-pdf") == 0
    assert len(gates(ws)) == before, "delivery added a gate receipt"
    delivery = [json.loads(l) for l
                in (ws / "journal.jsonl").read_text(encoding="utf-8").splitlines()
                if l.strip() and json.loads(l).get("action") == "delivery"][0]
    assert "verdict" not in delivery and "receipt_hash" not in delivery


def test_a_refused_pdf_is_named_in_the_record(tmp_path, monkeypatch):
    """The reason a PDF is missing has to survive the session that produced it.

    Asserted over the WHOLE list, never `[0]`. How many PDFs are refused depends
    on the machine: with a LaTeX engine only the CJK one is (its font is mocked
    away), without one every .md is, and `mock/assessment-1.md` sorts first. The
    first version of this test indexed [0], passed here and failed in CI — an
    environment-dependent assertion, which is a test bug and not a finding.
    """
    import json
    ws = build(tmp_path, md="# 岗位候选\n\n中文内容。\n")
    monkeypatch.setattr(deliver, "_render_portable_report",
                        lambda *args: (False, "disabled for this regression"))
    monkeypatch.setattr(deliver, "pick_cjk_font", lambda *args: None)
    assert run(ws, tmp_path / "out") == 2
    rec = [json.loads(l) for l in
           (ws / "journal.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
    delivery = [r for r in rec if r.get("action") == "delivery"][0]
    refused = delivery["pdf_refused"]
    assert refused, "a refused PDF left no trace"
    shortlist = [n for n in refused if "求职建议报告.pdf" in n]
    assert shortlist, f"the CJK document is not named among {refused}"
    assert "CJK font" in shortlist[0], (
        "the entry must carry WHY, not just that something was refused: "
        f"{shortlist[0]}")


def test_an_unwritable_workspace_journal_does_not_fail_the_delivery(tmp_path, monkeypatch):
    """The files are already copied at that point. Failing here would report a
    delivery that did happen as one that did not."""
    ws = build(tmp_path)
    monkeypatch.setattr(deliver.journal, "append",
                        lambda *a, **k: (_ for _ in ()).throw(OSError("read-only")))
    assert run(ws, tmp_path / "out", "--no-pdf") == 0


@pytest.mark.parametrize('text', ['You have an 80% chance of getting an interview.',
                                  'You are likely to be hired.'])
def test_delivery_refuses_predictions_in_client_report(tmp_path, text):
    ws = build(tmp_path, md=text)
    dest = tmp_path / 'out'
    assert run(ws, dest, '--no-pdf') == 2
    assert not dest.exists()


# Real discovery fixtures make this a delivery/evidence regression rather than
# a test that trusts a manually set "complete" flag.
def discovery_delivery(tmp_path, *, detail_only=False):
    import discover_fixtures as fx
    import yaml
    ws = fx.build_workspace(tmp_path)
    shortlist = yaml.safe_load((ws / "shortlist.yaml").read_text())
    if detail_only:
        shortlist["rows"] = [row for row in shortlist["rows"] if row["quality"] == "complete"]
        (ws / "shortlist.yaml").write_text(yaml.safe_dump(shortlist, allow_unicode=True))
        matches = yaml.safe_load((ws / "candidate-match.yaml").read_text())
        matches["rows"] = [row for row in matches["rows"] if row["basis"] == "detail"]
        (ws / "candidate-match.yaml").write_text(yaml.safe_dump(matches, allow_unicode=True))
    (ws / "report.md").write_text("# Career report\n\n" + "\n".join(
        f"[Posting]({row['url']})" for row in shortlist["rows"]), encoding="utf-8")
    deliver.journal.receipt(ws, "check_shortlist", {
        "shortlist.yaml": deliver.journal.sha256_file(ws / "shortlist.yaml"),
        "brief.yaml": deliver.journal.sha256_file(ws / "brief.yaml"),
        "journal.jsonl": deliver.journal.sha256_file(ws / "journal.jsonl"),
    }, "pass")
    from test_search_coverage import seed_coverage
    seed_coverage(ws)
    return ws, shortlist


def test_complete_discovery_refuses_unattempted_cards_even_after_search_passes(tmp_path):
    ws, _ = discovery_delivery(tmp_path)
    assert "full description not reviewed" in deliver._discover_handoff_problem(ws)
    assert run(ws, tmp_path / "out", "--no-pdf") == 2
    assert not (tmp_path / "out").exists()


def test_complete_discovery_accepts_full_evidence_and_checks_it_again(tmp_path):
    ws, _ = discovery_delivery(tmp_path, detail_only=True)
    assert deliver._discover_handoff_problem(ws) == ""
    assert run(ws, tmp_path / "out", "--no-pdf") == 0
    (ws / "raw/51job-detail-173199597.json").write_text("{}")
    assert "detail evidence is incomplete" in deliver._discover_handoff_problem(ws)


def test_preliminary_delivery_needs_an_explicit_request(tmp_path):
    import yaml
    ws, _ = discovery_delivery(tmp_path)
    brief = yaml.safe_load((ws / "brief.yaml").read_text())
    brief["report_scope"] = "preliminary"
    (ws / "brief.yaml").write_text(yaml.safe_dump(brief))
    deliver.journal.receipt(ws, "check_shortlist", {}, "pass")
    assert "explicit request" in deliver._discover_handoff_problem(ws)
    brief["preliminary_request"] = "Give me the initial list now; review details later."
    (ws / "brief.yaml").write_text(yaml.safe_dump(brief))
    assert deliver._discover_handoff_problem(ws) == ""


def test_complete_delivery_does_not_accept_a_stale_search_receipt(tmp_path):
    ws, _ = discovery_delivery(tmp_path, detail_only=True)
    with (ws / "shortlist.yaml").open("a") as handle:
        handle.write("\nchanged: true\n")
    assert "receipt is stale" in deliver._discover_handoff_problem(ws)


def test_multi_round_delivery_checks_each_retained_id_and_its_details(tmp_path):
    import yaml
    first, rows_a = discovery_delivery(tmp_path / "a", detail_only=True)
    second, rows_b = discovery_delivery(tmp_path / "b")
    ws = tmp_path / "collection"
    ws.mkdir()
    (ws / "report.md").write_text((first / "report.md").read_text() + (second / "report.md").read_text())
    manifest = {"rounds": [{"workspace": str(first)}, {"workspace": str(second)}]}
    (ws / "collection.yaml").write_text(yaml.safe_dump(manifest))
    from test_search_coverage import seed_coverage
    seed_coverage(ws, [first, second])
    assert "full description not reviewed" in deliver._discover_handoff_problem(ws)
    # A preliminary source round does not downgrade a complete collection.
    brief = yaml.safe_load((second / "brief.yaml").read_text())
    brief.update(report_scope="preliminary", preliminary_request="Show an early shortlist")
    (second / "brief.yaml").write_text(yaml.safe_dump(brief))
    deliver.journal.receipt(second, "check_shortlist", {}, "pass")
    assert "full description not reviewed" in deliver._discover_handoff_problem(ws)
    manifest["rounds"][1]["ids"] = [r["id"] for r in rows_b["rows"] if r["quality"] == "complete"]
    (ws / "collection.yaml").write_text(yaml.safe_dump(manifest))
    assert deliver._discover_handoff_problem(ws) == ""
    manifest["rounds"][1]["ids"] = ["invented-posting"]
    (ws / "collection.yaml").write_text(yaml.safe_dump(manifest))
    assert "existing shortlist rows" in deliver._discover_handoff_problem(ws)


def test_access_failure_exception_requires_a_real_capture_and_visible_reason(tmp_path):
    import yaml
    ws, shortlist = discovery_delivery(tmp_path)
    row = next(r for r in shortlist["rows"] if r["quality"] != "complete")
    reason = "The employer page timed out; full requirements could not be checked."
    name = "raw/51job-detail-failed.json"
    shortlist["detail_unavailable"] = [{"id": row["id"], "reason": reason, "capture": name}]
    (ws / "shortlist.yaml").write_text(yaml.safe_dump(shortlist, allow_unicode=True))
    deliver.journal.receipt(ws, "check_shortlist", {}, "pass")
    assert "not reviewed" in deliver._discover_handoff_problem(ws)
    (ws / name).write_text("request timeout")
    deliver.journal.append(ws, {"action": "adapter_call", "site": "51job", "command": "detail",
        "ts": "2026-08-09T14:04:00Z",
        "classification": "transport", "exit_code": 1, "stdout_file": name,
        "command_line": f"opencli 51job detail {row['source_id']}"})
    assert "not reviewed" in deliver._discover_handoff_problem(ws)
    with (ws / "report.md").open("a") as handle:
        handle.write("\n\n" + reason)
    assert deliver._discover_handoff_problem(ws) == ""


def test_site_stop_does_not_require_retrying_every_retained_posting(tmp_path):
    ws, shortlist = discovery_delivery(tmp_path)
    row = next(r for r in shortlist["rows"] if r["quality"] != "complete")
    name = "raw/site-login.json"
    (ws / name).write_text("login required")
    deliver.journal.append(ws, {"action": "adapter_call", "site": "51job", "command": "search",
        "classification": "not_logged_in", "exit_code": 1, "stdout_file": name})
    exception = {"reason": "Login required; details remain unverified.", "capture": name}
    assert deliver._unavailable_detail_supported(ws, row, exception, exception["reason"])
    row["source_site"] = "another-employer"
    assert not deliver._unavailable_detail_supported(ws, row, exception, exception["reason"])


# ---- several workspaces, one client folder ---------------------------------
#
# MEASURED (review, 2026-10-24): workspace acme (cv.md + cv.docx + report) and
# then workspace beta delivered to one --to both exited 0; the first report was
# gone and 简历/ held Acme's 简历.docx beside Beta's 简历.md and 简历.pdf — a mixed
# application the client could have sent. One workspace delivering apply and
# then interview into its default folder lost the apply report the same way.

def _pdf(path, text):
    import pymupdf
    with pymupdf.open() as doc:
        doc.new_page().insert_text((72, 72), text)
        doc.save(path)


def posting_workspace(root, name, company, role, marker, *, mode="apply", docx=False):
    ws = root / name
    ws.mkdir(parents=True)
    (ws / "report.md").write_text(f"# Advice\n\n{marker} report: confirm the scope first.\n",
                                  encoding="utf-8")
    (ws / "cv.md").write_text(f"# CV\n\n{marker} curriculum\n", encoding="utf-8")
    _pdf(ws / "cv.pdf", f"{marker} curriculum")
    if docx:
        from docx import Document
        document = Document()
        document.add_paragraph(f"{marker} curriculum")
        document.save(ws / "cv.docx")
    (ws / "posting.yaml").write_text(f'company: "{company}"\nrole_title: "{role}"\n',
                                     encoding="utf-8")
    deliver.journal.append(ws, {"action": "mode_entry", "mode": mode})
    return ws


def snapshot(dest):
    return {rel: (dest / rel).read_bytes() for rel in client_files(dest)}


def test_a_second_workspace_never_overwrites_the_first_in_a_shared_folder(tmp_path, capsys):
    acme = posting_workspace(tmp_path, "acme-ml-engineer-2026-10-24", "Acme BV",
                             "ML Engineer", "ACME", docx=True)
    beta = posting_workspace(tmp_path, "beta-data-scientist-2026-10-24", "Beta Corp",
                             "Data Scientist", "BETA")
    dest = tmp_path / "consultation"
    assert run(acme, dest) == 0
    first = snapshot(dest)
    assert set(first) == {"报告/Acme BV－ML Engineer－2026-10-24－投递建议报告.md", "报告/Acme BV－ML Engineer－2026-10-24－投递建议报告.pdf",
                          "简历/简历.md", "简历/简历.pdf", "简历/简历.docx"}
    capsys.readouterr()

    assert run(beta, dest) == 0
    output = capsys.readouterr()
    sub = "Beta Corp Data Scientist"
    assert {rel: data for rel, data in snapshot(dest).items() if not rel.startswith(sub)} == first
    assert client_files(dest / sub) == {"报告/Beta Corp－Data Scientist－2026-10-24－投递建议报告.md", "报告/Beta Corp－Data Scientist－2026-10-24－投递建议报告.pdf",
                                        "简历/简历.md", "简历/简历.pdf"}
    assert "BETA" in (dest / sub / "报告/Beta Corp－Data Scientist－2026-10-24－投递建议报告.md").read_text(encoding="utf-8")
    assert (dest / sub / "简历/简历.pdf").read_bytes() == (beta / "cv.pdf").read_bytes()
    assert "NOTICE_SEPARATE_FOLDER" in output.err and sub in output.err
    assert output.out.rstrip().endswith(str(dest / sub)), "the folder to quote is the real one"


def test_the_second_workspace_redelivers_into_its_own_subfolder(tmp_path):
    acme = posting_workspace(tmp_path, "acme", "Acme BV", "ML Engineer", "ACME")
    beta = posting_workspace(tmp_path, "beta", "Beta Corp", "Data Scientist", "BETA")
    dest = tmp_path / "consultation"
    assert run(acme, dest, "--no-pdf") == 0
    assert run(beta, dest, "--no-pdf") == 0
    (beta / "report.md").write_text("# Advice\n\nBETA round two.\n", encoding="utf-8")
    assert run(beta, dest, "--no-pdf") == 0
    folders = {p.name for p in dest.iterdir() if p.is_dir()} - {"报告", "简历"}
    assert folders == {"Beta Corp Data Scientist"}
    sub = dest / "Beta Corp Data Scientist"
    assert "round two" in (sub / "报告/Beta Corp－Data Scientist－投递建议报告.md").read_text(encoding="utf-8")
    assert "ACME" in (dest / "报告/Acme BV－ML Engineer－投递建议报告.md").read_text(encoding="utf-8")


def test_one_workspace_in_a_second_mode_keeps_the_first_modes_report(tmp_path, monkeypatch):
    monkeypatch.setattr(deliver, "DEFAULT_ROOT", tmp_path / "Downloads")
    ws = posting_workspace(tmp_path, "acme", "Acme BV", "ML Engineer", "APPLY")
    assert deliver.main(["--workspace", str(ws), "--no-pdf"]) == 0
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "interview"})
    (ws / "report.md").write_text("# Debrief\n\nINTERVIEW debrief.\n", encoding="utf-8")
    assert deliver.main(["--workspace", str(ws), "--no-pdf"]) == 0
    folder = tmp_path / "Downloads" / ws.name
    assert "APPLY" in (folder / "报告/Acme BV－ML Engineer－投递建议报告.md").read_text(encoding="utf-8")
    assert "INTERVIEW" in (folder / "模拟面试/报告/Acme BV－ML Engineer－面试复盘报告.md").read_text(encoding="utf-8")
    assert client_files(folder / "模拟面试") == {"报告/Acme BV－ML Engineer－面试复盘报告.md"}


def test_a_redelivery_removes_only_its_own_stale_files(tmp_path):
    ws = posting_workspace(tmp_path, "acme", "Acme BV", "ML Engineer", "ACME")
    (ws / "letter.md").write_text("Dear Acme,\n", encoding="utf-8")
    dest = tmp_path / "consultation"
    assert run(ws, dest, "--no-pdf") == 0
    assert (dest / "简历/求职信.md").is_file()
    (dest / "简历/notes.txt").write_text("the client's own note", encoding="utf-8")
    (ws / "letter.md").unlink()
    assert run(ws, dest, "--no-pdf") == 0
    assert not (dest / "简历/求职信.md").exists(), "a letter this round no longer has"
    assert (dest / "简历/notes.txt").read_text(encoding="utf-8") == "the client's own note"
    assert client_files(dest) == {"报告/Acme BV－ML Engineer－投递建议报告.md", "简历/简历.md", "简历/简历.pdf",
                                  "简历/notes.txt"}


def test_a_stale_file_the_client_edited_is_kept_and_named(tmp_path, capsys):
    ws = posting_workspace(tmp_path, "acme", "Acme BV", "ML Engineer", "ACME")
    (ws / "letter.md").write_text("Dear Acme,\n", encoding="utf-8")
    dest = tmp_path / "consultation"
    assert run(ws, dest, "--no-pdf") == 0
    (dest / "简历/求职信.md").write_text("Dear Acme, (edited by the client)\n", encoding="utf-8")
    (ws / "letter.md").unlink()
    capsys.readouterr()
    assert run(ws, dest, "--no-pdf") == 0
    assert "edited by the client" in (dest / "简历/求职信.md").read_text(encoding="utf-8")
    assert "NOTICE_STALE_FILE_KEPT" in capsys.readouterr().err


@pytest.mark.parametrize("manifest", [
    "{broken",
    '{"deliveries": "not a list"}',
    # A manifest must never steer the stale-file cleanup outside client files.
    '{"deliveries": [{"workspace": "WS", "mode": "unknown", "folder": "", '
    '"files": {"../victim.txt": "0"}}]}',
    '{"deliveries": [{"workspace": "WS", "mode": "unknown", "folder": "", '
    '"files": {"简历/../../victim.txt": "0"}}]}',
])
def test_an_untrustworthy_manifest_stops_delivery_before_copying(tmp_path, capsys, manifest):
    ws = build(tmp_path)
    dest = tmp_path / "consultation"
    dest.mkdir()
    (dest / MANIFEST).write_text(manifest.replace("WS", str(ws.resolve())),
                                              encoding="utf-8")
    (tmp_path / "victim.txt").write_text("keep", encoding="utf-8")
    assert run(ws, dest, "--no-pdf") == 2
    err = capsys.readouterr().err
    assert "DELIVER_MANIFEST_UNREADABLE" in err and "Traceback" not in err
    assert client_files(dest) == set()
    assert (tmp_path / "victim.txt").read_text(encoding="utf-8") == "keep"


# ---- client files an older deliver.py left with no manifest ------------------
#
# MEASURED (2026-10-24): before 1.3.0 deliver.py wrote the same 报告/ and 简历/
# names straight into the destination and recorded nothing. A folder holding
# such a delivery looked empty to the ownership check, and the next workspace
# overwrote its report and CV while exiting 0.

def _legacy_delivery(dest, marker):
    """What an older deliver.py left: client files, and no manifest."""
    (dest / "报告").mkdir(parents=True)
    (dest / "简历").mkdir()
    (dest / "报告/求职建议报告.md").write_text(f"# Advice\n\n{marker} report.\n", encoding="utf-8")
    _pdf(dest / "简历/简历.pdf", f"{marker} curriculum")
    return snapshot(dest)


def test_files_an_older_deliver_left_without_a_manifest_are_not_overwritten(tmp_path, capsys):
    beta = posting_workspace(tmp_path, "beta", "Beta Corp", "Data Scientist", "BETA")
    dest = tmp_path / "consultation"
    legacy = _legacy_delivery(dest, "ACME")
    assert not (dest / MANIFEST).exists()
    capsys.readouterr()
    assert run(beta, dest, "--no-pdf") == 0
    output = capsys.readouterr()
    sub = "Beta Corp Data Scientist"
    assert {rel: data for rel, data in snapshot(dest).items() if not rel.startswith(sub)} == legacy
    assert client_files(dest / sub) == {"报告/Beta Corp－Data Scientist－投递建议报告.md", "简历/简历.md", "简历/简历.pdf"}
    assert "NOTICE_SEPARATE_FOLDER" in output.err and sub in output.err
    assert "简历/简历.pdf" in output.err, "the notice names the files it stepped around"
    assert output.out.rstrip().endswith(str(dest / sub))


def test_the_next_workspace_does_not_overwrite_them_either(tmp_path):
    """The first delivery writes a manifest that records only its own subfolder.
    The unrecorded top-level files must stay protected for every later one."""
    beta = posting_workspace(tmp_path, "beta", "Beta Corp", "Data Scientist", "BETA")
    gamma = posting_workspace(tmp_path, "gamma", "Gamma AG", "ML Engineer", "GAMMA")
    dest = tmp_path / "consultation"
    legacy = _legacy_delivery(dest, "ACME")
    assert run(beta, dest, "--no-pdf") == 0
    assert run(gamma, dest, "--no-pdf") == 0
    after = snapshot(dest)
    assert {rel: after[rel] for rel in legacy} == legacy
    assert "GAMMA" in (dest / "Gamma AG ML Engineer/报告/Gamma AG－ML Engineer－投递建议报告.md").read_text(encoding="utf-8")
    assert run(beta, dest, "--no-pdf") == 0, "a redelivery returns to its own subfolder"
    assert {p.name for p in dest.iterdir() if p.is_dir()} == {
        "报告", "简历", "Beta Corp Data Scientist", "Gamma AG ML Engineer"}


def test_unrecorded_files_identical_to_this_rounds_copies_are_its_own(tmp_path, capsys):
    """The twin: an older deliver.py's copy of THIS workspace's files, byte for
    byte. Nothing would be lost by writing over it, so the round stays put."""
    ws = posting_workspace(tmp_path, "acme", "Acme BV", "ML Engineer", "ACME")
    dest = tmp_path / "consultation"
    (dest / "报告").mkdir(parents=True)
    (dest / "简历").mkdir()
    shutil.copy2(ws / "report.md", dest / "报告/求职建议报告.md")
    shutil.copy2(ws / "cv.pdf", dest / "简历/简历.pdf")
    capsys.readouterr()
    assert run(ws, dest, "--no-pdf") == 0
    assert "NOTICE_SEPARATE_FOLDER" not in capsys.readouterr().err
    assert client_files(dest) == {"报告/求职建议报告.md", "报告/Acme BV－ML Engineer－投递建议报告.md", "简历/简历.md", "简历/简历.pdf"}


def test_one_changed_byte_makes_the_unrecorded_files_foreign(tmp_path):
    ws = posting_workspace(tmp_path, "acme", "Acme BV", "ML Engineer", "ACME")
    dest = tmp_path / "consultation"
    (dest / "报告").mkdir(parents=True)
    (dest / "简历").mkdir()
    shutil.copy2(ws / "report.md", dest / "报告/求职建议报告.md")
    (dest / "简历/简历.pdf").write_bytes((ws / "cv.pdf").read_bytes() + b"\n")
    assert run(ws, dest, "--no-pdf") == 0
    assert (dest / "简历/简历.pdf").read_bytes() == (ws / "cv.pdf").read_bytes() + b"\n"
    assert client_files(dest / "Acme BV ML Engineer") == {
        "报告/Acme BV－ML Engineer－投递建议报告.md", "简历/简历.md", "简历/简历.pdf"}


def test_unrecorded_files_in_this_workspaces_default_folder_are_its_own(tmp_path, monkeypatch):
    """The other twin: ~/Downloads/<workspace name>/ is named after this very
    workspace, so an older deliver.py's files there are this workspace's earlier
    round, and the new one replaces them in place."""
    monkeypatch.setattr(deliver, "DEFAULT_ROOT", tmp_path / "Downloads")
    ws = posting_workspace(tmp_path, "acme-ml-engineer-2026-10-24", "Acme BV", "ML Engineer",
                           "ROUND-TWO")
    folder = tmp_path / "Downloads" / ws.name
    _legacy_delivery(folder, "ROUND-ONE")
    assert deliver.main(["--workspace", str(ws), "--no-pdf"]) == 0
    assert client_files(folder) == {"报告/求职建议报告.md", "报告/Acme BV－ML Engineer－2026-10-24－投递建议报告.md", "简历/简历.md", "简历/简历.pdf"}
    assert "ROUND-ONE" in (folder / "报告/求职建议报告.md").read_text(encoding="utf-8")
    assert "ROUND-TWO" in (folder / "报告/Acme BV－ML Engineer－2026-10-24－投递建议报告.md").read_text(encoding="utf-8")
    assert (folder / "简历/简历.pdf").read_bytes() == (ws / "cv.pdf").read_bytes()


def test_another_workspaces_default_folder_is_not_this_ones(tmp_path, monkeypatch):
    """Delivering INTO a folder named after a different workspace is a shared
    --to like any other: its unrecorded files are that workspace's."""
    monkeypatch.setattr(deliver, "DEFAULT_ROOT", tmp_path / "Downloads")
    beta = posting_workspace(tmp_path, "beta", "Beta Corp", "Data Scientist", "BETA")
    other = tmp_path / "Downloads" / "acme-ml-engineer-2026-10-24"
    legacy = _legacy_delivery(other, "ACME")
    assert run(beta, other, "--no-pdf") == 0
    assert {rel: data for rel, data in snapshot(other).items()
            if not rel.startswith("Beta Corp")} == legacy


def test_only_the_clients_own_files_leave_the_folder_usable(tmp_path):
    """Files deliver.py never writes are not a delivery: a client's own notes in
    the folder do not push the round into a subfolder."""
    ws = posting_workspace(tmp_path, "acme", "Acme BV", "ML Engineer", "ACME")
    dest = tmp_path / "consultation"
    (dest / "简历").mkdir(parents=True)
    (dest / "简历/notes.txt").write_text("the client's own note", encoding="utf-8")
    (dest / "求职建议报告.md").write_text("a draft the client saved here", encoding="utf-8")
    assert run(ws, dest, "--no-pdf") == 0
    assert client_files(dest) == {"报告/Acme BV－ML Engineer－投递建议报告.md", "简历/简历.md", "简历/简历.pdf",
                                  "简历/notes.txt", "求职建议报告.md"}


# ---- CV and letter PDFs come from their renderers, never from the report one --
#
# MEASURED: with no cv.pdf, cv.md went through the REPORT renderer into 简历.pdf
# — footer "Career consultation report / Page 1", raw Markdown markers, extra
# "Link:" lines — bypassing render_cv's NO_ENGINE/TEXT_OFF_PAGE refusals and
# check_pages, and the delivery reported itself complete.

def test_a_cv_without_its_own_pdf_ships_as_text_and_says_why(tmp_path, capsys):
    from docx import Document
    ws = tmp_path / "acme"
    ws.mkdir()
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "apply"})
    (ws / "report.md").write_text("# Advice\n\nSend the tailored CV.\n", encoding="utf-8")
    (ws / "cv.md").write_text("# Jane Doe\n\n**Engineer** at [Acme](https://acme.example)\n",
                              encoding="utf-8")
    document = Document()
    document.add_paragraph("Jane Doe")
    document.save(ws / "cv.docx")
    (ws / "letter.md").write_text("Dear Acme,\n", encoding="utf-8")
    before = {p.name for p in ws.iterdir()}
    dest = tmp_path / "out"
    assert run(ws, dest) == 2, "a missing CV PDF is an incomplete apply delivery"
    err = capsys.readouterr().err
    assert client_files(dest) == {"报告/求职建议报告.md", "报告/求职建议报告.pdf",
                                  "简历/简历.md", "简历/简历.docx", "简历/求职信.md"}
    assert "NOTICE_PDF_MISSING: 简历.pdf" in err and "cv.pdf" in err and "render_cv.py" in err
    assert "NOTICE_PDF_MISSING: 求职信.pdf" in err and "render_letter.py" in err
    # Only the report's own render record may appear; none may claim a CV PDF.
    assert {p.name for p in ws.iterdir()} == before | {"report-render.json"}


def test_a_supporting_statement_ships_as_text_with_a_notice(tmp_path, capsys):
    ws = tmp_path / "uni"
    ws.mkdir()
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "apply"})
    (ws / "report.md").write_text("# Advice\n\nSubmit the statement in the portal.\n",
                                  encoding="utf-8")
    (ws / "supporting-statement.md").write_text("I meet criterion 1 because ...\n",
                                                encoding="utf-8")
    dest = tmp_path / "out"
    assert run(ws, dest) == 0
    assert client_files(dest) == {"报告/求职建议报告.md", "报告/求职建议报告.pdf",
                                  "简历/申请陈述.md"}
    assert "NOTICE_PDF_NOT_PRODUCED: 申请陈述.pdf" in capsys.readouterr().err


def test_an_explicit_copy_outside_apply_names_a_missing_cv_pdf_without_failing(tmp_path, capsys):
    ws = tmp_path / "interview"
    ws.mkdir()
    deliver.journal.append(ws, {"action": "mode_entry", "mode": "interview"})
    (ws / "report.md").write_text("# Debrief\n\nPractise the system design answer.\n",
                                  encoding="utf-8")
    (ws / "cv.md").write_text("# Jane Doe\n", encoding="utf-8")
    dest = tmp_path / "out"
    assert run(ws, dest, "--include-applications") == 0
    assert client_files(dest) == {"报告/求职建议报告.md", "报告/求职建议报告.pdf",
                                  "简历/简历.md"}
    assert "NOTICE_PDF_NOT_PRODUCED: 简历.pdf" in capsys.readouterr().err


# ---- could not run is exit 2 with one line, never a traceback ---------------

def test_an_unreadable_journal_is_named_as_such(tmp_path, capsys):
    ws = tmp_path / "ws"
    ws.mkdir()
    (ws / "report.md").write_text("# Advice\n", encoding="utf-8")
    (ws / "journal.jsonl").mkdir()
    assert run(ws, tmp_path / "out", "--no-pdf") == 2
    err = capsys.readouterr().err
    assert err.startswith("DELIVER_WORKSPACE_UNREADABLE:") and "Traceback" not in err


def test_an_unknown_option_is_a_usage_error_not_a_traceback(tmp_path):
    ws = build(tmp_path)
    r = subprocess.run([sys.executable, str(REPO / "scripts" / "deliver.py"),
                        "--workspace", str(ws), "--files", str(tmp_path / "missing.md")],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 2 and "Traceback" not in r.stderr and "--files" in r.stderr
