import copy
import pathlib
import sys

import pymupdf
import pytest
import yaml
from docx import Document

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import check_apply
import check_layout
import journal
import deliver


def reviewed_workspace(tmp_path):
    ws = tmp_path / "application"
    ws.mkdir()
    docx = Document()
    docx.add_paragraph("Test Candidate")
    docx.save(ws / "cv.docx")
    reference = tmp_path / "template.docx"
    docx.save(reference)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Test Candidate")
    doc.save(ws / "cv.pdf")
    page.get_pixmap().save(ws / "preview.png")
    doc.close()
    requirements = {
        "reference": {"path": str(reference), "sha256": journal.sha256_file(reference)},
        "expectations": {key: "Measured reference: " + key for key in check_layout.CHECKS},
        "page_size_pt": [595, 842], "text_bounds_pt": [72, 59, 523, 770],
        "docx_geometry_twips": {"width": 12240, "height": 15840, "top": 1440,
                                "bottom": 1440, "left": 1800, "right": 1800},
        "text_samples": [{"role": role, "text": "Test Candidate", "size_pt": 11,
                          "fonts": ["Helvetica"]} for role in ("name", "heading", "body")],
    }
    (ws / "layout-requirements.yaml").write_text(yaml.safe_dump(requirements))
    (ws / "word-export.pdf").write_bytes((ws / "cv.pdf").read_bytes())
    review = {
        "requirements_sha256": journal.sha256_file(ws / "layout-requirements.yaml"),
        "word_export": {"path": "word-export.pdf", "sha256": journal.sha256_file(ws / "cv.pdf"),
                        "docx_sha256": journal.sha256_file(ws / "cv.docx")},
        "reviewer": "Test reviewer", "reviewed_at": "2026-09-11T12:00:00+08:00",
        "reference": {"path": str(reference), "sha256": journal.sha256_file(reference)},
        "user_overrides": [],
        "checks": {key: {"status": "pass", "detail": "Compared to the reference."}
                   for key in check_layout.CHECKS},
        "unresolved_differences": [],
        "artifacts": {name: {"sha256": journal.sha256_file(ws / name), "pages": 1,
                             "previews": [{"page": 1, "path": "preview.png",
                                           "sha256": journal.sha256_file(ws / "preview.png")}]}
                      for name in ("cv.docx", "cv.pdf")},
    }
    write_review(ws, review)
    return ws, review


def write_review(ws, review):
    (ws / "layout-review.yaml").write_text(yaml.safe_dump(review), encoding="utf-8")


def test_absent_workspace_is_not_created(tmp_path):
    ws = tmp_path / "missing"
    assert check_layout.main(["--workspace", str(ws)]) == 2
    assert not ws.exists()


def test_rendered_cv_requires_layout_review_in_apply(tmp_path):
    ws, _ = reviewed_workspace(tmp_path)
    assert ("check_layout", "a rendered CV exists") in check_apply.conditional_gates(ws)
    (ws / "layout-review.yaml").unlink()
    assert check_layout.main(["--workspace", str(ws)]) == 1
    assert "NO_LAYOUT_REVIEW" in journal.read_receipts(ws, "check_layout")[-1]["findings"][0]


def test_complete_visual_record_is_bound_to_actual_files(tmp_path):
    ws, _ = reviewed_workspace(tmp_path)
    assert check_layout.main(["--workspace", str(ws)]) == 0
    receipt = journal.read_receipts(ws, "check_layout")[-1]
    assert receipt["input_hashes"]["cv.docx"] == journal.sha256_file(ws / "cv.docx")
    assert receipt["input_hashes"]["cv.pdf"] == journal.sha256_file(ws / "cv.pdf")
    assert journal.receipt_intact(receipt)
    assert check_apply._stale_inputs(ws, "check_layout", receipt) == []


def test_layout_receipt_cannot_bind_an_unrelated_external_file(tmp_path):
    ws, _ = reviewed_workspace(tmp_path)
    assert check_layout.main(["--workspace", str(ws)]) == 0
    receipt = journal.read_receipts(ws, "check_layout")[-1]
    unrelated = tmp_path / "unrelated.txt"
    unrelated.write_text("not the supplied template")
    receipt["input_hashes"][str(unrelated)] = journal.sha256_file(unrelated)
    findings = check_apply._stale_inputs(ws, "check_layout", receipt)
    assert any("RECEIPT_INPUT_OUTSIDE_WORKSPACE" in item for item in findings)


@pytest.mark.parametrize("changed", ["cv.docx", "cv.pdf", "preview.png", "reference"])
def test_changed_output_preview_or_template_invalidates_review(tmp_path, changed):
    ws, review = reviewed_workspace(tmp_path)
    target = pathlib.Path(review["reference"]["path"]) if changed == "reference" else ws / changed
    target.write_bytes(target.read_bytes() + b"changed")
    assert check_layout.main(["--workspace", str(ws)]) == 1
    assert any("STALE_LAYOUT_REVIEW" in s for s in journal.read_receipts(ws, "check_layout")[-1]["findings"])


@pytest.mark.parametrize("problem", ["missing_page", "duplicate_page", "wrong_count", "failed_check", "blank_check", "unresolved", "not_image"])
def test_unreviewed_or_unresolved_layout_cannot_pass(tmp_path, problem):
    ws, original = reviewed_workspace(tmp_path)
    review = copy.deepcopy(original)
    pdf = review["artifacts"]["cv.pdf"]
    if problem == "missing_page":
        pdf["previews"] = []
    elif problem == "duplicate_page":
        pdf["previews"] *= 2
    elif problem == "wrong_count":
        pdf["pages"] = 2
    elif problem == "failed_check":
        review["checks"]["page_geometry"]["status"] = "fail"
    elif problem == "blank_check":
        review["checks"]["fonts"]["detail"] = ""
    elif problem == "unresolved":
        review["unresolved_differences"] = ["Unexpected top margin."]
    elif problem == "not_image":
        (ws / "preview.png").write_text("not a rendered page")
        for item in review["artifacts"].values():
            item["previews"][0]["sha256"] = journal.sha256_file(ws / "preview.png")
    write_review(ws, review)
    assert check_layout.main(["--workspace", str(ws)]) == 1


def test_no_template_does_not_waive_page_review(tmp_path):
    ws, review = reviewed_workspace(tmp_path)
    review["reference"] = None
    requirements = yaml.safe_load((ws / "layout-requirements.yaml").read_text())
    requirements["reference"] = None
    (ws / "layout-requirements.yaml").write_text(yaml.safe_dump(requirements))
    review["requirements_sha256"] = journal.sha256_file(ws / "layout-requirements.yaml")
    write_review(ws, review)
    assert check_layout.main(["--workspace", str(ws)]) == 0
    review["artifacts"]["cv.pdf"]["previews"] = []
    write_review(ws, review)
    assert check_layout.main(["--workspace", str(ws)]) == 1


@pytest.mark.parametrize("problem", ["size", "font", "margin", "requirements_missing", "requirements_stale", "word_export"])
def test_measured_format_failures_cannot_be_waived_by_passed_visual_strings(tmp_path, problem):
    ws, review = reviewed_workspace(tmp_path)
    if problem.startswith("requirements"):
        path = ws / "layout-requirements.yaml"
        if problem == "requirements_missing":
            path.unlink()
        else:
            path.write_text(path.read_text() + "\nchanged: true\n")
    elif problem == "margin":
        document = Document(ws / "cv.docx")
        document.sections[0].left_margin = 1000000
        document.save(ws / "cv.docx")
        review["artifacts"]["cv.docx"]["sha256"] = journal.sha256_file(ws / "cv.docx")
        review["word_export"]["docx_sha256"] = journal.sha256_file(ws / "cv.docx")
    else:
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Test Candidate", fontsize=14 if problem == "size" else 11,
                         fontname="cour" if problem == "font" else "helv")
        doc.save(ws / "new.pdf")
        doc.close()
        (ws / "new.pdf").replace(ws / "cv.pdf")
        review["artifacts"]["cv.pdf"]["sha256"] = journal.sha256_file(ws / "cv.pdf")
        if problem != "word_export":
            (ws / "word-export.pdf").write_bytes((ws / "cv.pdf").read_bytes())
            review["word_export"]["sha256"] = journal.sha256_file(ws / "cv.pdf")
    write_review(ws, review)
    assert check_layout.main(["--workspace", str(ws)]) == 1
    findings = journal.read_receipts(ws, "check_layout")[-1]["findings"]
    code = ("MISSING_LAYOUT_INPUT" if problem == "requirements_missing" else
            "STALE_LAYOUT_REVIEW" if problem == "requirements_stale" else
            "PDF_NOT_WORD_EXPORT" if problem == "word_export" else "FORMAT_MISMATCH")
    assert any(code in item for item in findings)


def reviewed_report(tmp_path):
    ws, review = reviewed_workspace(tmp_path)
    requirements = yaml.safe_load((ws / "layout-requirements.yaml").read_text())
    requirements["text_samples"][0]["role"] = "title"
    (ws / "report-layout-requirements.yaml").write_text(yaml.safe_dump(requirements))
    (ws / "report.md").write_text("# Test Candidate\n")
    (ws / "report.pdf").write_bytes((ws / "cv.pdf").read_bytes())
    review["requirements_sha256"] = journal.sha256_file(ws / "report-layout-requirements.yaml")
    review["source_sha256"] = journal.sha256_file(ws / "report.md")
    review["artifacts"] = {"report.pdf": review["artifacts"]["cv.pdf"]}
    (ws / "report-layout-review.yaml").write_text(yaml.safe_dump(review))
    return ws


def test_full_page_requirement_rejects_large_blank_bottom(tmp_path):
    ws, review = reviewed_workspace(tmp_path)
    path = ws / "layout-requirements.yaml"
    requirements = yaml.safe_load(path.read_text())
    requirements["minimum_content_bottom_pt"] = 777
    path.write_text(yaml.safe_dump(requirements))
    review["requirements_sha256"] = journal.sha256_file(path)
    write_review(ws, review)
    assert check_layout.main(["--workspace", str(ws)]) == 1
    assert any("excessive lower-page whitespace" in s for s in journal.read_receipts(ws, "check_layout")[-1]["findings"])


@pytest.mark.parametrize("text", ["本次围绕固定虚构履历进行了11个问答，并非真人面试。",
                                 "This fixed fictional profile is a synthetic candidate, not a real interview."])
def test_client_delivery_rejects_internal_test_narration(tmp_path, text):
    ws = tmp_path / "application"
    ws.mkdir()
    (ws / "report.md").write_text(text)
    written, problems = deliver.deliver(ws, tmp_path / "delivery", "test", make_pdf=False)
    assert written == []
    assert any("REPORT_INTERNAL_NARRATION" in s for s in problems)


def test_audience_check_keeps_normal_testing_and_interview_advice():
    assert deliver.report_audience_findings("软件测试工程师要求JUnit测试经验。模拟面试前，整理测试输入、预期结果和缺陷复测记录。") == []


def test_delivery_preserves_reviewed_pdf_without_rerendering(tmp_path, monkeypatch):
    ws = reviewed_report(tmp_path)
    assert check_layout.main(["--workspace", str(ws), "--report"]) == 0
    def forbidden(*args, **kwargs):
        pytest.fail("must not regenerate a reviewed report")
    monkeypatch.setattr(deliver, "render_pdf", forbidden)
    written, problems = deliver.deliver(ws, tmp_path / "delivery", "test", include_applications=False)
    assert not problems
    pdf = next(path for path in written if path.suffix == ".pdf")
    assert pdf.read_bytes() == (ws / "report.pdf").read_bytes()


@pytest.mark.parametrize("changed", ["report.md", "report.pdf", "report-layout-requirements.yaml", "report-layout-review.yaml"])
def test_delivery_blocks_stale_or_removed_report_review(tmp_path, changed):
    ws = reviewed_report(tmp_path)
    target = ws / changed
    if changed == "report-layout-review.yaml":
        target.unlink()
    else:
        target.write_bytes(target.read_bytes() + b"changed")
    written, problems = deliver.deliver(ws, tmp_path / "delivery", "test", include_applications=False)
    assert written == [] and problems
    assert not (tmp_path / "delivery").exists()


# 2026-09-19 review: a span's box includes its whitespace. LibreOffice exports
# wrapped lines with a trailing space, so the skill's own English example CV
# "overflowed" by 1.5 pt of blank, while invisible NBSP padding "filled" a page.
def replace_pdf(ws, review, lines, *, requirements=None):
    doc = pymupdf.open()
    page = doc.new_page()
    for origin, text in [((72, 72), "Test Candidate"), *lines]:
        page.insert_text(origin, text, fontsize=11, fontname="helv")
    doc.save(ws / "new.pdf")
    doc.close()
    (ws / "new.pdf").replace(ws / "cv.pdf")
    (ws / "word-export.pdf").write_bytes((ws / "cv.pdf").read_bytes())
    review["artifacts"]["cv.pdf"]["sha256"] = journal.sha256_file(ws / "cv.pdf")
    review["word_export"]["sha256"] = journal.sha256_file(ws / "cv.pdf")
    if requirements:
        path = ws / "layout-requirements.yaml"
        data = yaml.safe_load(path.read_text())
        data.update(requirements)
        path.write_text(yaml.safe_dump(data))
        review["requirements_sha256"] = journal.sha256_file(path)
    write_review(ws, review)


def layout_findings(ws):
    check_layout.main(["--workspace", str(ws)])
    return journal.read_receipts(ws, "check_layout")[-1]["findings"]


def test_trailing_spaces_past_the_margin_are_not_painted_text(tmp_path):
    ws, review = reviewed_workspace(tmp_path)
    # Glyphs end at x=516.9; the three trailing spaces extend the span to 526.1.
    replace_pdf(ws, review, [((350, 120), "Compared six competing products   ")])
    assert not [f for f in layout_findings(ws) if "outside required bounds" in f]


def test_painted_text_past_the_margin_still_fails(tmp_path):
    ws, review = reviewed_workspace(tmp_path)
    replace_pdf(ws, review, [((400, 120), "Compared six competing products")])
    assert [f for f in layout_findings(ws) if "outside required bounds" in f]


def test_invisible_padding_does_not_fill_the_page(tmp_path):
    ws, review = reviewed_workspace(tmp_path)
    replace_pdf(ws, review, [((72, 765), "  ")], requirements={"minimum_content_bottom_pt": 760})
    assert [f for f in layout_findings(ws) if "excessive lower-page whitespace" in f]


def test_printed_content_at_the_bottom_fills_the_page(tmp_path):
    ws, review = reviewed_workspace(tmp_path)
    replace_pdf(ws, review, [((72, 765), "References available")], requirements={"minimum_content_bottom_pt": 760})
    assert not [f for f in layout_findings(ws) if "excessive lower-page whitespace" in f]


# 2026-09-19 review: interview mode runs one round per dispatch, so a real
# client's report naturally says what this round covered and that a mock is not
# a real interview; a chemistry posting can ask for synthetic candidate routes.
@pytest.mark.parametrize("text", ["本轮只练了技术面，行为面下次再练。",
                                  "这只是练习，并非真人面试，请把反馈当作准备建议。",
                                  "This mock is not a real interview; treat the scores as practice notes.",
                                  "The role designs synthetic candidate routes for new APIs."])
def test_audience_check_delivers_real_practice_and_posting_language(text):
    assert deliver.report_audience_findings(text) == []


# 2026-09-19 review: the reviewed report.pdf was bound to report.md only by
# hashes the agent writes. Re-hashing an edited report.md without re-rendering
# delivered a PDF saying "Apply to Acme now" beside Markdown saying "Do not apply".
def reviewed_report_with(tmp_path, markdown, lines):
    ws = reviewed_report(tmp_path)
    (ws / "report.md").write_text(markdown, encoding="utf-8")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="cjk", fontbuffer=pymupdf.Font("china-s").buffer)
    page.insert_text((72, 72), "Test Candidate", fontsize=11, fontname="helv")
    for y, text in lines:
        page.insert_text((72, y), text, fontsize=11,
                         fontname="cjk" if deliver.has_cjk(text) else "helv")
    doc.save(ws / "new.pdf")
    doc.close()
    (ws / "new.pdf").replace(ws / "report.pdf")
    review = yaml.safe_load((ws / "report-layout-review.yaml").read_text())
    review["source_sha256"] = journal.sha256_file(ws / "report.md")
    review["artifacts"]["report.pdf"]["sha256"] = journal.sha256_file(ws / "report.pdf")
    (ws / "report-layout-review.yaml").write_text(yaml.safe_dump(review))
    return ws


def test_a_reviewed_report_pdf_must_carry_the_current_markdown(tmp_path):
    ws = reviewed_report_with(tmp_path, "# Test Candidate\n\nDo not apply to Acme: German C1 is required.\n",
                              [(100, "Apply to Acme now.")])
    findings, _ = check_layout.inspect(ws, report=True)
    assert any("REPORT_PDF_TEXT_MISMATCH" in f and "Do not apply" in f for f in findings)
    written, problems = deliver.deliver(ws, tmp_path / "out", "test")
    assert written == [] and any("REPORT_PDF_TEXT_MISMATCH" in p for p in problems)


def test_a_rerendered_report_pdf_passes_the_text_binding(tmp_path):
    ws = reviewed_report_with(tmp_path, "# Test Candidate\n\nDo not apply to Acme: German C1 is required.\n",
                              [(100, "Do not apply to Acme: German C1 is"), (114, "required.")])
    findings, _ = check_layout.inspect(ws, report=True)
    assert not [f for f in findings if "REPORT_PDF_TEXT_MISMATCH" in f]


def test_merged_company_cells_in_a_reviewed_report_are_delivered(tmp_path):
    company, roles = "中信建投证券股份有限公司", ["投行分析师", "债券分析师", "并购分析师",
                                          "保荐承销岗", "财务顾问岗", "资产证券化岗"]
    markdown = "# Test Candidate\n\n| 公司 | 岗位 | 地点 |\n|---|---|---|\n" + "".join(
        f"| {company} | {role} | 上海 |\n" for role in roles)
    # The PDF names the company once in a merged cell, as report-writing.md prefers.
    lines = [(100, "公司 岗位 地点"), (114, f"{company} {roles[0]} 上海")] + [
        (128 + 14 * i, f"{role} 上海") for i, role in enumerate(roles[1:])]
    ws = reviewed_report_with(tmp_path, markdown, lines)
    written, problems = deliver.deliver(ws, tmp_path / "out", "test")
    assert any(p.parent.name == "报告" and p.suffix == ".pdf" for p in written), problems


# MuPDF's built-in fonts draw a visible "\u00b7" for these, so the rule is tested
# on characters; the effect was checked on real Chrome-rendered filler PDFs.
@pytest.mark.parametrize("char", ["\u2800", "\u3164", "\u200b", "\ufeff", "\uffa0", "\u115f", "\u00a0", " "])
def test_invisible_filler_characters_are_not_ink(char):
    import layout_requirements
    assert not layout_requirements.is_ink(char)


@pytest.mark.parametrize("char", ["A", "\u4e2d", "\u00b7", "\u2022", "_", "-", "1"])
def test_printed_characters_are_ink(char):
    import layout_requirements
    assert layout_requirements.is_ink(char)


@pytest.mark.parametrize("text", ["本次使用固定虚构简历完成演练。", "这份测试用的虚构履历只用于演示。",
                                  "本次围绕固定虚构\n履历进行了问答。", "This fixed fictional\nprofile was used.",
                                  "A fixed fictional CV was used for this run."])
def test_fixture_narration_variants_are_refused(text):
    assert deliver.report_audience_findings(text)


def test_fictional_test_data_advice_for_a_qa_role_is_delivered():
    assert deliver.report_audience_findings("面试前准备测试用的虚构数据集，并说明边界用例的设计。") == []


# 2026-09-19 adversarial pass on the report text binding. Page texts mimic
# PyMuPDF extraction: one string per page, one PDF line per "\n".
def binding(markdown, *pages, tables=None):
    import layout_requirements
    return layout_requirements.report_text_problems(markdown, list(pages))


@pytest.mark.parametrize("markdown,pdf", [
    ("Apply to Acme Imaging now.\n", "Do not apply to Acme Imaging now.\n"),
    ("建议投递第一个岗位。\n", "不建议投递第一个岗位。\n"),
    ("Apply to Acme now.\n", "Apply to Acme now, but only after German C1.\n"),
])
def test_text_removed_from_the_markdown_is_caught(markdown, pdf):
    assert binding(markdown, pdf)


TABLE = "| 优先级 | 岗位 | 薪资 |\n|---|---|---|\n| 1 | 投行分析师 | 30-50K |\n| 2 | 债券分析师 | 20-30K |\n| 3 | 并购分析师 | 25-40K |\n"
TABLE_PDF = "优先级 岗位 薪资\n1 投行分析师 30-50K\n2 债券分析师 20-30K\n3 并购分析师 25-40K\n"
TABLE_ROWS = [[["优先级", "岗位", "薪资"], ["1", "投行分析师", "30-50K"],
               ["2", "债券分析师", "20-30K"], ["3", "并购分析师", "25-40K"]]]


def test_an_honest_table_binds():
    assert binding(TABLE, TABLE_PDF, tables=TABLE_ROWS) == []


@pytest.mark.parametrize("edit", [
    ("| 1 | 投行分析师 |", "| 1 | 债券分析师 |"),  # priorities swapped...
    ("| 3 | 并购分析师 | 25-40K |", "| 3 | 并购分析师 | 30-50K |"),  # value copied from another row
])
@pytest.mark.xfail(strict=True, reason="table cells are checked for presence, not placement")
def test_a_table_edit_using_values_from_other_rows_is_caught(edit):
    changed = TABLE.replace(*edit)
    if edit[1] == "| 1 | 债券分析师 |":
        changed = changed.replace("| 2 | 债券分析师 |", "| 2 | 投行分析师 |")
    assert binding(changed, TABLE_PDF, tables=TABLE_ROWS)


def test_a_paragraph_across_a_page_break_with_furniture_binds():
    markdown = "## Summary\n\nThe Leiden role matches your reconstruction thesis and the team hires graduates each spring.\n"
    page1 = "Career report\nSummary\nThe Leiden role matches your reconstruction\nPage 1 of 2\n"
    page2 = "Career report\nthesis and the team hires graduates each spring.\nPage 2 of 2\n"
    assert binding(markdown, page1, page2) == []


def test_a_table_cell_across_a_page_break_with_a_repeated_header_binds():
    markdown = "| 公司 | 岗位 | 地点 |\n|---|---|---|\n| 中信建投证券股份有限公司 | 投资银行部股权业务线分析师 | 上海 |\n"
    page1 = "公司 岗位 地点\n中信建投证券股份有限公司 投资银行部股权\n第 1 页\n"
    page2 = "公司 岗位 地点\n业务线分析师 上海\n第 2 页\n"
    assert binding(markdown, page1, page2) == []


def test_common_markdown_forms_bind():
    markdown = ("1. First step\n1. Second step\n1. Third step\n\n"
                "Read the [posting][p] before applying.\n\n[p]: https://example.com/jobs/12345\n\n"
                "Salary &amp; benefits&nbsp;are listed.\n\n"
                "The team uses PyTorch[^tools] daily.\n\n[^tools]: Per the posting.\n")
    pdf = ("1. First step\n2. Second step\n3. Third step\nRead the posting before applying.\n"
           "Salary & benefits are listed.\nThe team uses PyTorch1 daily.\n1 Per the posting.\n")
    assert binding(markdown, pdf) == []


def test_card_layouts_may_omit_table_header_labels():
    markdown = "| 优先级 | 公司与地点 |\n|---|---|\n| 1 | 联影医疗·上海 |\n"
    assert binding(markdown, "1\n联影医疗·上海\n") == []


# 2026-09-19 second adversarial pass on the report text binding.
CELLS = "| 公司 | 签证 | 结论 |\n|---|---|---|\n| ANWB | 否 | 暂不投递 |\n| 联影 | 是 | 建议投递 |\n"
CELLS_PDF = "公司 签证 结论\nANWB 否 暂不投递\n联影 是 建议投递\n"
CELLS_ROWS = [[["公司", "签证", "结论"], ["ANWB", "否", "暂不投递"], ["联影", "是", "建议投递"]]]


# Known gap, pinned so that strengthening the check has to update it: a table
# cell is only required to be present somewhere, because wrapped and CJK
# columns leave no reliable reading of where a cell stands.
@pytest.mark.xfail(strict=True, reason="table cells are checked for presence, not placement")
def test_a_verdict_cell_rewritten_inside_the_old_text_is_caught():
    assert binding(CELLS.replace("| ANWB | 否 | 暂不投递 |", "| ANWB | 否 | 投递 |"), CELLS_PDF, tables=CELLS_ROWS)


def test_a_changed_header_label_is_caught():
    assert binding(CELLS.replace("| 公司 | 签证 | 结论 |", "| 公司 | 签证 | 初判 |"), CELLS_PDF, tables=CELLS_ROWS)


def test_an_honest_table_with_headers_binds():
    assert binding(CELLS, CELLS_PDF, tables=CELLS_ROWS) == []


@pytest.mark.parametrize("markdown,pdf", [
    ("### 3. 大模型后训练算法专家\n", "1. 大模型后训练算法专家\n"),
    ("- 4 of the 25 roles require Dutch.\n", "- 14 of the 25 roles require Dutch.\n"),
    ("- Visa sponsorship at ANWB is confirmed.\n", "- No visa sponsorship at ANWB is confirmed.\n"),
])
def test_edits_hidden_beside_existing_text_are_caught(markdown, pdf):
    assert binding(markdown, pdf)


def test_a_two_character_cell_across_a_page_break_binds():
    markdown = "| 岗位 | 薪资 |\n|---|---|\n| 视觉slam算法 | 35-55K x 14薪 |\n"
    page1 = "岗位 薪资\n视觉\n第 1 页\n"
    page2 = "岗位 薪资\nslam算法 35-55K x 14薪\n第 2 页\n"
    assert binding(markdown, page1, page2) == []


def test_pandoc_and_css_authoring_forms_bind():
    markdown = ("---\ntitle: 中国 AI 岗位检索报告\nauthor: Consultant\n---\n\n"
                "## 结论 {#top-picks .unnumbered}\n\n建议优先投递联影。\n\n\\newpage\n\n"
                "<style>ol { list-style-type: cjk-ideographic; }</style>\n\n"
                "1. 先投联影\n1. 再投迈瑞\n")
    pdf = "中国 AI 岗位检索报告\nConsultant\n结论\n建议优先投递联影。\n一、 先投联影\n二、 再投迈瑞\n"
    assert binding(markdown, pdf) == []


def test_a_sentence_deleted_next_to_a_page_break_is_caught():
    markdown = "荷兰岗位的月薪区间需要自行核对。请在面试前确认合同细节。\n"
    page1 = "求职建议报告\n荷兰岗位的月薪区间需要自行核对。按荷兰法定最低假期工资比例折算，月基本工资约 3086 至 3858 欧元。\n第 1 页\n"
    page2 = "求职建议报告\n请在面试前确认合同细节。\n第 2 页\n"
    assert binding(markdown, page1, page2)

