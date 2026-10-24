#!/usr/bin/env python3
"""Render a motivation letter (letter.yaml) to Markdown / .docx / PDF.

Usage:
  python render_letter.py LETTER.yaml --format md|docx|pdf --out OUTPATH
"""
import argparse
import pathlib
import re
# The compile itself now runs inside render_cv.compile_latex, but this import
# stays: `render_letter.subprocess` is the same module object as
# `render_cv.subprocess`, and it is the handle the letter's engine-argv test
# patches to prove the two renderers build the same argv.
import subprocess  # noqa: F401
import sys


# Reuse LaTeX helpers from render_cv (same directory).
sys.path.insert(0, str(pathlib.Path(__file__).parent))
import journal
import render_cv
from docx_content import DocxMarkupError, save_clean_docx


def load(path):
    """Raises journal.YamlUnreadable; main() prints it and exits 2. Same reasoning as
    render_cv.load_profile — no workspace, so no receipt, and the exit code is the
    only channel left. 2 stays "could not read the input", 1 stays "the PDF failed"."""
    return journal.load_yaml(path)


def body_paragraphs(d):
    body = d.get("body") or []
    return [body] if isinstance(body, str) else body


def render_markdown(d):
    s = d.get("sender", {}) or {}
    r = d.get("recipient", {}) or {}
    lines = [s.get("name", ""), s.get("email", ""), s.get("location", ""), ""]
    if d.get("date"):
        lines += [d["date"], ""]
    lines += [r.get("name", ""), r.get("company", ""), r.get("location", ""), ""]
    if salutation_for(d):
        lines += [salutation_for(d), ""]
    for para in body_paragraphs(d):
        lines += [para, ""]
    lines += [x for x in (closing_for(d), s.get("name", "")) if x]
    return "\n".join([ln for ln in lines]) + "\n"


def render_docx(d, out_path):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml.ns import qn
    doc = Document()
    meta = _meta(d)
    section = doc.sections[0]
    section.page_width, section.page_height = (
        (Mm(215.9), Mm(279.4)) if render_cv.paper_for(meta) == "letterpaper"
        else (Mm(210), Mm(297)))
    section.top_margin = section.bottom_margin = Mm(25)
    section.left_margin = section.right_margin = Mm(25)
    style = doc.styles["Normal"]
    style.font.name = meta.get("main_font") or "Times New Roman"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)
    east_asia = meta.get("cjk_font") or {
        "zh": "SimSun", "ja": "Yu Mincho", "ko": "Malgun Gothic"
    }.get(meta["language"])
    if east_asia:
        fonts.set(qn("w:eastAsia"), east_asia)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15
    s = d.get("sender", {}) or {}
    r = d.get("recipient", {}) or {}
    for line in [s.get("name"), s.get("email"), s.get("location")]:
        if line:
            doc.add_paragraph(line)
    if d.get("date"):
        doc.add_paragraph(d["date"])
    for line in [r.get("name"), r.get("company"), r.get("location")]:
        if line:
            doc.add_paragraph(line)
    if salutation_for(d):
        doc.add_paragraph(salutation_for(d))
    for para in body_paragraphs(d):
        doc.add_paragraph(para)
    if closing_for(d):
        doc.add_paragraph(closing_for(d))
    doc.add_paragraph(s.get("name", ""))
    save_clean_docx(doc, out_path)


# Salutations for an UNNAMED recipient, quoted from references/motivation-letter.md
# :208, which is the file that sources them. Not invented here and not translated
# here: that line is the provenance, and a second spelling in this file would be a
# second thing for the two to disagree about.
#
# Every default was `"Dear Hiring Manager,"` / `"Sincerely,"` regardless of
# language, so a Dutch letter shipped with an English opener and an English
# sign-off — against motivation-letter.md:272, "Never mix languages in one
# letter", in the renderer that file drives.
_SALUTATION = {
    "en": "Dear Hiring Team,",
    "de": "Sehr geehrte Damen und Herren,",
    "fr": "Madame, Monsieur,",
    "es": "Estimados señores:",
    "it": "Gentile Responsabile delle Assunzioni,",
    "nl": "Geachte heer/mevrouw,",
}
# Quoted from references/motivation-letter.md, the same file and the same status
# as the salutations above. An independent pass caught the asymmetry: five
# languages had a sourced opener and no sourced sign-off, so an honest Dutch or
# German letter that left `closing` unset — exactly as leaving `salutation`
# unset correctly works — hard-failed NO_CLOSING with nothing to fill it in.
#
# Languages still absent here (zh, ja, ko and the rest) keep the honest
# behaviour: no sign-off is emitted and check_letter reports the gap, because
# inventing one is what this skill bans.
_CLOSING = {
    "en": "Sincerely,",
    "de": "Mit freundlichen Grüßen",
    "nl": "Met vriendelijke groet,",
    "fr": "Cordialement,",
    "es": "Atentamente,",
    "it": "Cordiali saluti,",
}


# ── what language the letter is in, when it does not say ─────────────────────
#
# references/motivation-letter.md §4 and modes/apply.md give letter.yaml an
# "exact schema" with no `meta` block, and this file read `meta.language` and
# `meta.target_market` with English and A4 as the defaults. So a letter written
# to the documented schema could not be anything else: a Chinese letter with no
# salutation rendered "Dear Hiring Team," and "Sincerely," around Chinese
# paragraphs (and check_letter passed it), and a US application shipped cv.pdf on
# Letter beside letter.pdf and letter.docx on A4.
#
# The body is the evidence. Script settles CJK outright (see `infer_language` for
# why CJK characters are weighed against Latin words). For Latin script, a short
# list of function words per language is counted, and a language is named only
# when it clearly wins; a body that does not say stays undecided, rather than
# being guessed, and the old default applies.
_KANA = ("\u3040", "\u30ff")
_HANGUL = (("\u1100", "\u11ff"), ("\u3130", "\u318f"), ("\uac00", "\ud7af"))
_HAN = (("\u3400", "\u4dbf"), ("\u4e00", "\u9fff"), ("\uf900", "\ufaff"))


def _in(ch, ranges) -> bool:
    return any(lo <= ch <= hi for lo, hi in ranges)


def script_counts(text) -> dict:
    """{'kana', 'hangul', 'han', 'latin'} character counts."""
    counts = {"kana": 0, "hangul": 0, "han": 0, "latin": 0}
    for ch in str(text or ""):
        if _KANA[0] <= ch <= _KANA[1]:
            counts["kana"] += 1
        elif _in(ch, _HANGUL):
            counts["hangul"] += 1
        elif _in(ch, _HAN):
            counts["han"] += 1
        elif ch.isalpha() and (ch.isascii() or "\u00c0" <= ch <= "\u024f"):
            counts["latin"] += 1
    return counts


def cjk_count(counts) -> int:
    return counts["kana"] + counts["hangul"] + counts["han"]


# Function words, chosen for being frequent in a formal letter and hard to meet
# in another of these languages: a word that is ordinary in two of them ("de",
# "la", "in", "is", "was", "es", "do") is left out of both, so one shared word
# cannot tip the count. `_LATIN_MIN_SHARE` and the two-to-one margin are why a
# short or mixed body comes back undecided.
#
# Left out on purpose, because they are ordinary in a language they do not
# belong to: en am/will (German), de die (Dutch) / des (French), nl de/en
# (French, Spanish), fr le (Italian), it il (French), es del/con/mi/su (Italian)
# and para/por/este (Portuguese), pt um (German) / dos (Spanish) / sua (Italian).
_FUNCTION_WORDS = {
    "en": "the and of to that with for my have this your at be which from by it "
          "are we our would not but has been their you",
    "de": "der das und ich nicht mit für ist ein eine einen zu den dem von auf bei "
          "ihre ihr sich auch dass mich mein meine habe wir sie über durch sehr "
          "ihnen bin",
    "nl": "het een van ik niet voor op bij dat zijn uw mijn ook naar om heb wij "
          "graag kan wordt deze aan maar mij",
    "fr": "les et une je j qu pour avec dans est mon mes votre vos au sur pas ce "
          "cette nous vous aux suis très par plus été",
    "es": "el los las y mis sus más pero soy muy nuestro nuestra también cuando "
          "estoy hasta",
    "it": "gli di della delle degli dei dalla che è mio mia sono nel nella alla ho "
          "anche questo questa più essere vostro vostra presso",
    "pt": "os uma em na é meu minha seu ao mas sou muito nossa nosso também não eu",
}
_WORD_SETS = {lang: set(words.split()) for lang, words in _FUNCTION_WORDS.items()}
_SHARED = {w for lang, ws in _WORD_SETS.items() for w in ws
           if any(w in other for o, other in _WORD_SETS.items() if o != lang)}
_WORD_SETS = {lang: ws - _SHARED for lang, ws in _WORD_SETS.items()}
_LATIN_MIN_SHARE = 0.08
_LATIN_WORD_RE = re.compile(r"[a-z\u00df-\u00f6\u00f8-\u024f]+")


def infer_language(text):
    """'zh' | 'ja' | 'ko' | 'en' | 'de' | 'nl' | 'fr' | 'es' | 'it' | 'pt', or None
    when the text does not settle it.

    CJK characters are weighed against Latin WORDS, not Latin letters. One CJK
    character carries about what a short word does, and a Chinese letter keeps
    team and tool names in English — "我在Google Cloud Platform的Distributed
    Systems团队负责…" has more Latin letters than Han and is still Chinese."""
    counts = script_counts(text)
    cjk = cjk_count(counts)
    words = _LATIN_WORD_RE.findall(str(text or "").lower())
    if cjk and cjk >= len(words):
        if counts["hangul"] > counts["han"] + counts["kana"]:
            return "ko"
        if counts["kana"] >= 0.1 * (counts["han"] + counts["kana"]):
            return "ja"
        return "zh"
    if not words:
        return None
    scores = sorted(((sum(w in ws for w in words), lang)
                     for lang, ws in _WORD_SETS.items()), reverse=True)
    (best, lang), (second, _) = scores[0], scores[1]
    if best >= 3 and best >= _LATIN_MIN_SHARE * len(words) and best >= 2 * second:
        return lang
    return None


def body_text(d) -> str:
    return " ".join(str(p) for p in body_paragraphs(journal.as_mapping(d)))


def declared_language(d) -> str:
    meta = journal.as_mapping(journal.as_mapping(d).get("meta"))
    return str(meta.get("language") or "").strip().lower()[:2]


def _language(d) -> str:
    """meta.language, else the body's own language, else English (the old default)."""
    return declared_language(d) or infer_language(body_text(d)) or "en"


def _meta(d) -> dict:
    """`meta` with `language` filled in, for the font and paper choices."""
    meta = dict(journal.as_mapping(journal.as_mapping(d).get("meta")))
    meta["language"] = _language(d)
    return meta


# ── the rest of the application package ──────────────────────────────────────
#
# The letter lives in the application workspace beside tailored-profile.yaml,
# whose `meta.target_market` the user chose in Step 0 and which already decides
# the CV's paper. Reading it is what puts the two documents of one package on one
# paper size. posting.yaml has no market field, only the posting's own location
# text, so it is consulted last and only to recognise a US/Canada posting.
_WORKSPACE_PROFILES = ("tailored-profile.yaml", "profile.yaml")
_INHERITED_FONTS = ("main_font", "cjk_font")


def _read_mapping(path):
    try:
        return journal.load_yaml(path, dict) if path.is_file() else {}
    except journal.YamlUnreadable:
        # The CV renderer and its gates report an unreadable profile; a letter
        # falls back to its own defaults rather than failing over a sibling file.
        return {}


def resolved_letter(d, workspace=None):
    """`(letter, notes)`: `d` with the meta it left out filled in from evidence.

    Each key `d` states itself wins. Otherwise: `language` from the body
    (`infer_language`), then from the workspace profile; `target_market`/`paper`
    from the workspace profile, then — for Letter paper only — from a US/Canada
    location in posting.yaml; `main_font`/`cjk_font` from the workspace profile,
    so a font the CV needed to keep its glyphs is the letter's font too. `notes`
    says what was taken from where, one line each, for the caller to print.
    """
    d = journal.as_mapping(d)
    own = journal.as_mapping(d.get("meta"))
    meta = dict(own)
    notes = []
    ws = pathlib.Path(workspace) if workspace is not None else None
    profile_meta, profile_name = {}, None
    if ws is not None:
        for name in _WORKSPACE_PROFILES:
            found = journal.as_mapping(_read_mapping(ws / name).get("meta"))
            if found:
                profile_meta, profile_name = found, name
                break

    if not (own.get("target_market") or own.get("paper")):
        taken = [k for k in ("target_market", "paper") if profile_meta.get(k)]
        for key in taken:
            meta[key] = profile_meta[key]
        if taken:
            notes.append(f"letter.yaml has no meta.target_market; using "
                         f"{', '.join(f'{k}={profile_meta[k]!r}' for k in taken)} "
                         f"from {profile_name}, so the letter prints on the CV's "
                         f"paper size")
        elif ws is not None:
            location = _read_mapping(ws / "posting.yaml").get("location")
            if isinstance(location, str) and render_cv.north_america(location):
                meta["paper"] = "letter"
                notes.append(f"letter.yaml has no meta.target_market; posting.yaml's "
                             f"location {location!r} is in the US/Canada, so the "
                             f"letter prints on Letter paper")
    for key in _INHERITED_FONTS:
        if not own.get(key) and profile_meta.get(key):
            meta[key] = profile_meta[key]

    if not declared_language(d):
        inferred = infer_language(body_text(d))
        fallback = str(profile_meta.get("language") or "").strip().lower()[:2]
        if inferred:
            meta["language"] = inferred
            notes.append(f"letter.yaml has no meta.language; the body reads as "
                         f"{inferred!r}")
        elif fallback:
            meta["language"] = fallback
            notes.append(f"letter.yaml has no meta.language and the body does not "
                         f"settle it; using {profile_name}'s {fallback!r}")
    out = dict(d)
    out["meta"] = meta
    return out, notes


def salutation_for(d) -> str:
    """The letter's own `salutation`, or the sourced default for its language.

    "" when there is no sourced default — Chinese, Japanese and Korean have
    none in motivation-letter.md, and defaulting them to `Dear Hiring Team,`
    is the exact defect this function was written to remove, just moved. The
    renderer omits the line and check_letter reports the gap.
    """
    written = str(d.get("salutation") or "").strip()
    if written:
        return written
    return _SALUTATION.get(_language(d), "")


def closing_for(d) -> str:
    """The letter's own `closing`, or the sourced default — "" when there is
    none for this language, because an English sign-off on a non-English letter
    is the defect, not the fallback."""
    written = str(d.get("closing") or "").strip()
    if written:
        return written
    return _CLOSING.get(_language(d), "")


def build_latex(d, engine=None, cjk=False):
    """Assemble the letter's LaTeX source.

    The preamble comes from `render_cv.latex_preamble`, not from a copy here.
    This file used to carry its own `inputenc`+`fontenc` pair, which meant it
    also carried its own copy of the bug that pair causes under the engine
    `find_latex_engine` actually prefers: a letter to `Ștefan Ionescu` at
    `Politechnika Śląska` compiled, exited 0, and reached him addressed to
    "tefan Ionescu" at "Politechnika lska". See `latex_preamble` for why the
    choice is made from the engine and not from the text.
    """
    e = render_cv.latex_escape
    s = d.get("sender", {}) or {}
    r = d.get("recipient", {}) or {}
    # The sender's name keeps a space between CJK characters, as the CV's name
    # does: `山田 太郎` is a Japanese name, not a folded scalar.
    sender_name = e(s.get("name", ""), keep_cjk_spaces=True)
    # `meta` is passed for the paper size. Without it the CV in an application
    # package printed on Letter for a US role while the letter beside it printed
    # on A4 — one package, two page sizes, and nothing said so. `_meta` adds the
    # letter's language, which picks the CJK font chain for a Japanese letter.
    parts = render_cv.latex_preamble(engine=engine, cjk=cjk, margin="2.5cm",
                                     meta=_meta(d)) + [
        r"\pagestyle{empty}",
        r"\begin{document}",
        r"\noindent %s\\ %s\\ %s\\[1em]" % (
            sender_name, e(s.get("email", "")), e(s.get("location", ""))),
    ]
    if d.get("date"):
        parts.append(r"%s\\[1em]" % e(d["date"]))
    parts.append(r"\noindent %s\\ %s\\ %s\\[1em]" % (
        e(r.get("name", "")), e(r.get("company", "")), e(r.get("location", ""))))
    if salutation_for(d):
        parts.append(e(salutation_for(d)) + r"\\[1em]")
    for para in body_paragraphs(d):
        parts.append(e(para) + r"\\[1em]")
    if closing_for(d):
        parts.append(e(closing_for(d)) + r"\\")
    parts.append(sender_name)
    parts.append(r"\end{document}")
    return "\n".join(parts) + "\n"


def render_pdf(d, out_path, reasons=None):
    out_path = pathlib.Path(out_path)
    tex_path = out_path.with_suffix(".tex")
    # Any PDF here is from an earlier run; this call either replaces it or must
    # leave none. Same reasoning, same helper, as render_cv.render_pdf — the two
    # early returns below (unsupported script, no engine) are exactly the paths
    # that used to leave a superseded letter sitting under the delivered name.
    render_cv._discard_pdf(out_path, tex_path)

    # Right-to-left first, and before any source is written. Same template, same
    # missing `bidi`, same consequence as render_cv: the compile succeeds and
    # produces a reversed, unshaped page at exit 0. The .tex is refused with it,
    # because it is the source that produces exactly that page.
    if render_cv.profile_has_rtl(d):
        print("WARNING: the letter contains right-to-left text (Arabic, Hebrew "
              "or similar). The LaTeX template has no bidi support, so a PDF "
              "built from it would come out reversed and unshaped. No PDF was "
              "written; Markdown and .docx carry the text correctly.",
              file=sys.stderr)
        return render_cv._note(reasons, render_cv.UNSUPPORTED_SCRIPT)

    # CJK is TYPESET, not refused. render_cv has had an xeCJK preamble and a
    # font chain for a long time; this file simply never passed `cjk=` through,
    # so a Chinese application shipped a Chinese CV as PDF and no letter PDF at
    # all — one package, one script, two different answers, exit 0 both times.
    cjk = render_cv._has_cjk(str(d))
    # Resolve the engine first: the preamble has to match whatever will compile it.
    engine = render_cv.find_latex_engine(cjk=cjk)
    tex = build_latex(d, engine=engine, cjk=cjk)
    tex_path.write_text(tex, encoding="utf-8")
    if cjk and engine is None:
        print("WARNING: no Unicode LaTeX engine found for this CJK letter. A CJK "
              "PDF needs XeLaTeX (xelatex) or tectonic plus an installed CJK "
              f"font; pdflatex cannot do it. Source at {tex_path}.",
              file=sys.stderr)
        return render_cv._note(reasons, render_cv.NO_ENGINE)
    if engine is None:
        print(f"WARNING: no LaTeX engine found. Wrote {tex_path}.", file=sys.stderr)
        return render_cv._note(reasons, render_cv.NO_ENGINE)
    # One compile helper, both renderers — including the `Missing character`
    # scan, so a letter can no longer ship with glyphs silently dropped either.
    return render_cv.compile_latex(engine, tex_path, out_path, reasons=reasons)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("letter")
    ap.add_argument("--format", choices=["md", "docx", "pdf"], default="md")
    ap.add_argument("--out", required=True)
    args = ap.parse_args(argv)
    try:
        d = load(args.letter)
    except journal.YamlUnreadable as exc:
        print(f"cannot render: {exc}", file=sys.stderr)
        return 2
    render_cv.reset_photo_warnings()
    out = pathlib.Path(args.out)
    # modes/apply.md renders `<workspace>/letter.yaml`, so the letter's own
    # directory IS the application workspace, and the tailored profile beside it
    # is what knows the market. Said on stderr, so nothing is silently assumed.
    d, notes = resolved_letter(d, workspace=pathlib.Path(args.letter).resolve().parent)
    for note in notes:
        print(f"NOTICE: {note}.", file=sys.stderr)

    # Same early validation as render_cv.main: the letter shares `paper_for`, so
    # an override it cannot read must fail the same way for every format rather
    # than only on the LaTeX path.
    try:
        render_cv.paper_for((d or {}).get("meta"))
    except ValueError as exc:
        print(f"cannot render: {exc}", file=sys.stderr)
        return 2
    if args.format == "md":
        out.write_text(render_markdown(d), encoding="utf-8")
    elif args.format == "docx":
        try:
            render_docx(d, out)
        except DocxMarkupError as exc:
            print(f"cannot render Word: {exc}", file=sys.stderr)
            return 1
    else:
        reasons = []
        ok = render_pdf(d, out, reasons=reasons)
        if ok:
            if not render_cv.confirm_written(out, "pdf"):
                return 1
            print(f"Wrote {out}")
            return 0
        print(f"PDF not built; LaTeX source at {out.with_suffix('.tex')}",
              file=sys.stderr)
        # Same split as render_cv.main, from the same table: a missing engine or
        # a script this template cannot set are documented degradations; anything
        # else means this machine could have produced a correct PDF and did not.
        return 0 if render_cv.pdf_failure_is_tolerated(reasons) else 1
    if not render_cv.confirm_written(out, args.format):
        return 1
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    from cli_io import configure_output

    configure_output()
    sys.exit(main())
