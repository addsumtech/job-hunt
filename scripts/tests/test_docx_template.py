"""Observable layout and content contracts of the reviewed Word template."""
import pathlib
import sys
import zipfile
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import render_cv

@pytest.mark.parametrize('language', ['en', 'zh', 'ja', 'ko', 'de', 'fr', 'es', 'it', 'nl'])
def test_document_layout_keeps_content_and_editable_rules(tmp_path, language):
    profile = {
        'meta': {'name': 'Example Candidate', 'language': language, 'target_market': 'other'},
        'contact': {'email': 'candidate@example.com'},
        'experience': [{'org': 'Example Company', 'title': 'Analyst', 'location': 'Test City',
                        'start': '2022', 'end': '2024',
                        'bullets': ['Analysis: Reviewed a deliberately long dataset description.']}],
        'education': [{'institution': 'Example College', 'degree': 'BSc',
                       'location': 'Test Country', 'start': '2018', 'end': '2022',
                       'details': 'Coursework: Databases and statistics.'}],
        'projects': [{'name': 'Example Project', 'description': 'A test tool',
                      'links': [{'label': 'Source', 'url': 'https://example.com/project'}]}],
        'skills': {'tools': ['SQL', 'Python']},
    }
    path = tmp_path / 'cv.docx'
    render_cv.render_docx(profile, path)
    doc = Document(path)
    assert doc.styles['Title'].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.styles['Heading 1'].font.color.rgb == (0, 0, 0)
    rule = doc.styles['Heading 1'].element.find('.//' + qn('w:bottom'))
    assert rule is not None and rule.get(qn('w:val')) == 'single'
    assert not doc.styles['Title'].element.findall('.//' + qn('w:pBdr'))
    assert not any(p._p.findall('.//' + qn('w:u')) for p in doc.paragraphs if p.style.name == 'Heading 1')
    org = next(p for p in doc.paragraphs if p.text.startswith('Example Company'))
    assert org.text.endswith('2022 – 2024') and org.runs[0].bold
    assert org.paragraph_format.tab_stops[0].alignment == WD_TAB_ALIGNMENT.RIGHT
    assert not org.runs[-1].bold
    assert next(p for p in doc.paragraphs if p.text.startswith('Analyst')).text == 'Analyst\tTest City'
    edu = next(p for p in doc.paragraphs if p.text.startswith('BSc'))
    assert '2018 – 2022' in edu.text
    labeled = next(p for p in doc.paragraphs if p.text.startswith('Analysis:'))
    assert labeled.runs[0].bold and not labeled.runs[1].bold
    coursework = next(p for p in doc.paragraphs if p.text.startswith('Coursework:'))
    assert not any(r.bold for r in coursework.runs)
    for p in doc.paragraphs:
        if p.style.name == 'List Bullet':
            assert p.paragraph_format.left_indent.pt == 17
            assert p.paragraph_format.first_line_indent.pt == -17
    with zipfile.ZipFile(path) as z:
        assert b'https://example.com/project' in z.read('word/_rels/document.xml.rels')
        assert not any(n.startswith('word/media/') for n in z.namelist())
