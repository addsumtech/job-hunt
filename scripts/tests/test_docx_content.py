import pathlib
import sys
import zipfile

import pytest
import yaml
from docx import Document
from docx.oxml import OxmlElement

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import check_layout
import deliver
import docx_content
import render_cv
import render_letter
import render_rirekisho


FORMULA = r"在 $8 \times \text{H200}$ 上完成 200k 步预训练（约 47 小时）与 100k 步微调。"
PLAIN = "在 8 × H200 上完成 200k 步预训练（约 47 小时）与 100k 步微调。"


def profile(text):
    return {"meta": {"name": "Example Candidate", "language": "zh", "target_market": "China"},
            "contact": {"email": "candidate@example.com"},
            "projects": [{"name": "smolVLA 基座预训练", "description": text}],
            "experience": [{"org": "Example", "title": "Engineer", "bullets": [text]}]}


@pytest.mark.parametrize("text", [FORMULA, "**模型训练**", "*model training*", "__模型训练__",
                                   "`model_name`", "[论文](https://example.com/paper)",
                                   "## 项目经历", "- 训练模型", "```python", "|---|---|",
                                   r"\(x^2\)", r"\[x+y\]", "$$x+y$$", "$x^2$",
                                   r"使用 \frac{a}{b}"])
def test_rejects_source_notation(text):
    assert docx_content.markup_findings(text)


@pytest.mark.parametrize("text", [PLAIN, "C# / C++ / A* / R&D", "model_name / __init__ / user_id",
                                   "Budget $100 and $200; HK$500; US$1,000.",
                                   "8 × H200; 47 h; 2% and 10%; n = 8; x²",
                                   r"C:\temp\text\times\model.py", "2 ** 3 = 8"])
def test_preserves_real_technical_text(text):
    assert docx_content.markup_findings(text) == []


@pytest.mark.parametrize("kind", ["cv", "letter", "rirekisho"])
def test_all_renderers_refuse_formula_and_render_repaired_text(tmp_path, kind):
    output = tmp_path / f"{kind}.docx"

    def render(text):
        if kind == "cv":
            return render_cv.render_docx(profile(text), output)
        if kind == "letter":
            return render_letter.render_docx({"sender": {"name": "Example"}, "body": [text]}, output)
        return render_rirekisho.render_docx({**profile(text), "jp": {"motivation": text}}, output)

    with pytest.raises(docx_content.DocxMarkupError, match="DOCX_MARKUP"):
        render(FORMULA)
    assert not output.exists()
    render(PLAIN)
    assert docx_content.docx_findings(output) == []
    with zipfile.ZipFile(output) as archive:
        assert PLAIN in archive.read("word/document.xml").decode()


@pytest.mark.parametrize("kind", ["cv", "letter", "rirekisho"])
def test_cli_failure_names_the_problem_without_claiming_success(tmp_path, capsys, kind):
    data = profile(FORMULA)
    if kind == "letter":
        data = {"sender": {"name": "Example"}, "body": [FORMULA]}
    elif kind == "rirekisho":
        data["jp"] = {"motivation": FORMULA}
    source, output = tmp_path / "input.yaml", tmp_path / "result.docx"
    source.write_text(yaml.safe_dump(data, allow_unicode=True))
    args = [str(source), "--format", "docx", "--out", str(output)]
    renderer = {"cv": render_cv, "letter": render_letter, "rirekisho": render_rirekisho}[kind]
    if kind == "rirekisho":
        args.append("--allow-blank-dates")
    assert renderer.main(args) == 1
    captured = capsys.readouterr()
    assert "DOCX_MARKUP" in captured.err and "Wrote" not in captured.out
    assert "Traceback" not in captured.err
    assert not output.exists()


@pytest.mark.parametrize("location", ["body", "table", "header", "footer", "text_box"])
def test_checks_final_word_parts_across_runs(tmp_path, location):
    doc = Document()
    if location == "body":
        paragraph = doc.add_paragraph()
    elif location == "table":
        paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    elif location in {"header", "footer"}:
        paragraph = getattr(doc.sections[0], location).paragraphs[0]
    else:
        from docx.text.paragraph import Paragraph
        box = OxmlElement("w:txbxContent")
        inner = OxmlElement("w:p")
        box.append(inner)
        doc.add_paragraph()._p.append(box)
        paragraph = Paragraph(inner, doc._body)
    for text in ("项目描述：$8 ", "\\ti", "mes ", "\\text{H200}$"):
        paragraph.add_run(text)
    path = tmp_path / "manual.docx"
    doc.save(path)
    problems = docx_content.docx_findings(path)
    assert any("LaTeX" in problem and "paragraph" in problem for problem in problems)


def test_native_word_formatting_and_equations_remain_allowed(tmp_path):
    doc = Document()
    doc.add_heading("项目经历", level=1)
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run("项目描述：").bold = True
    paragraph.add_run(PLAIN)
    render_cv._add_hyperlink(paragraph, "论文", "https://example.com/paper")
    equation = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x² + y²"
    run.append(text)
    equation.append(run)
    paragraph._p.append(equation)
    path = tmp_path / "native.docx"
    docx_content.save_clean_docx(doc, path)
    assert docx_content.docx_findings(path) == []
    final = Document(path)
    assert final.paragraphs[0].style.name == "Heading 1"
    assert final.paragraphs[1].style.name == "List Bullet"
    assert final.paragraphs[1].runs[0].bold
    assert final.paragraphs[1]._p.xpath(".//m:oMath")


@pytest.mark.parametrize("name", ["cv", "letter", "rirekisho", "supporting-statement", "report"])
def test_delivery_refuses_any_word_residue_before_copying_files(tmp_path, name):
    ws, dest = tmp_path / "workspace", tmp_path / "delivery"
    ws.mkdir()
    (ws / "cv.md").write_text("# Example\n")
    doc = Document()
    doc.add_paragraph(FORMULA)
    doc.save(ws / f"{name}.docx")
    written, notes = deliver.deliver(ws, dest, "example", make_pdf=False)
    assert written == [] and any("DOCX_MARKUP" in note for note in notes)
    assert not dest.exists()


def test_layout_review_does_not_hide_residue(tmp_path):
    doc = Document()
    doc.add_paragraph(FORMULA)
    doc.save(tmp_path / "cv.docx")
    findings, _ = check_layout.inspect(tmp_path)
    assert any("DOCX_MARKUP" in finding for finding in findings)


def test_corrupt_word_file_is_not_a_clean_scan(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_text("**not a Word file**")
    assert docx_content.docx_findings(path)[0].startswith("DOCX_UNREADABLE")
