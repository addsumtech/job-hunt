import pathlib
import sys

import pytest

sys.path.insert(0, str(pathlib.Path(__file__).parent.parent))
import render_rirekisho as rr


@pytest.fixture
def jp_profile():
    return {
        "meta": {"name": "山田 太郎", "language": "ja"},
        "contact": {"email": "taro@example.com", "phone": "090-1234-5678",
                    "location": "東京都"},
        "jp": {
            "name_furigana": "やまだ たろう",
            "date_of_birth": "1995-04-12", "age": "31", "gender": "男",
            "address": "東京都新宿区1-1-1", "address_furigana": "とうきょうと…",
            "date": "2026年6月6日",
            "motivation": "貴社の研究開発に貢献したいと考え志望しました。",
            "personal_request": "貴社の規定に従います。",
        },
        "experience": [
            {"title": "エンジニア", "org": "株式会社ABC",
             "start": "2020-04", "end": "2023-03", "bullets": ["開発を担当。"]},
            {"title": "研究員", "org": "DEF研究所",
             "start": "2023-04", "end": "present", "bullets": ["研究に従事。"]},
        ],
        "education": [
            {"degree": "情報工学 学士", "institution": "○○大学",
             "start": "2014-04", "end": "2018-03"},
        ],
        "certifications": ["基本情報技術者試験 合格"],
    }


def test_ym_parsing():
    assert rr._ym("2023-09") == ("2023", "9")
    assert rr._ym("2021") == ("2021", "")
    assert rr._ym("2023年9月") == ("2023", "9")
    assert rr._ym("present") == ("", "")


def test_gakureki_shokureki_structure(jp_profile):
    rows = rr.gakureki_shokureki_rows(jp_profile)
    texts = [r["text"] for r in rows]
    # Section headers and the closing marker are present and ordered.
    assert "学歴" in texts and "職歴" in texts and "以上" in texts
    assert texts.index("学歴") < texts.index("職歴") < texts.index("以上")
    # Education yields enrolled + graduated; current job yields 現在に至る.
    assert any("入学" in t for t in texts)
    assert any("卒業" in t for t in texts)
    assert any("株式会社ABC　入社" in t for t in texts)
    assert "現在に至る" in texts
    # Chronological: earliest education enrollment comes before the latest job.
    enroll = next(r for r in rows if "入学" in r["text"])
    assert enroll["y"] == "2014"


def test_render_markdown_contains_form_sections(jp_profile):
    md = rr.render_markdown(jp_profile)
    assert md.startswith("# 履歴書")
    assert "## 学歴・職歴" in md
    assert "## 免許・資格" in md
    assert "## 志望の動機" in md
    assert "山田 太郎" in md
    assert "やまだ たろう" in md          # furigana surfaced
    assert "| 年 | 月 | 学歴・職歴 |" in md  # table header


def test_render_docx_builds_form(jp_profile, tmp_path):
    from docx import Document
    out = tmp_path / "rirekisho.docx"
    rr.render_docx(jp_profile, out)
    assert out.exists()
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "履歴書" in text
    # The photo placeholder appears when no photo_path is given.
    all_cell_text = "\n".join(
        cell.text for tbl in doc.tables for row in tbl.rows for cell in row.cells)
    assert "写真貼付欄" in all_cell_text
    assert "現在に至る" in all_cell_text       # work-history row rendered
    assert "貴社の研究開発" in all_cell_text     # motivation box rendered


def test_no_certifications_shows_placeholder():
    profile = {"meta": {"name": "X"}, "experience": [], "education": [],
               "certifications": []}
    rows = rr.licenses_rows(profile)
    assert rows == [{"y": "", "m": "", "text": "特になし"}]


def test_missing_jp_block_renders_with_placeholders(tmp_path):
    """A profile with no jp: block still renders the form (photo placeholder, no crash)."""
    from docx import Document
    profile = {"meta": {"name": "山田 太郎"}, "contact": {"email": "t@x.com"},
               "experience": [{"title": "研究員", "org": "X研", "start": "2023-04", "end": "present"}],
               "education": [{"degree": "学士", "institution": "○○大学",
                              "start": "2014-04", "end": "2018-03"}]}
    md = rr.render_markdown(profile)
    assert "# 履歴書" in md
    out = tmp_path / "r.docx"
    rr.render_docx(profile, out)
    cells = "\n".join(c.text for t in Document(str(out)).tables for row in t.rows for c in row.cells)
    assert "写真貼付欄" in cells


ACCEPTED = ["2023-09", "2021", "2023年9月", "2023/09"]
REJECTED = ["Sep 2023", "September 2023", "03/2021", "", None]


@pytest.mark.parametrize("value", ACCEPTED)
def test_accepted_date_formats_produce_no_fatal_problem(jp_profile, value):
    """`warnings` is not asserted empty here: the jp_profile fixture's
    certification is genuinely undated (`"基本情報技術者試験 合格"`), and an
    undated certification is a warning by design."""
    p = dict(jp_profile)
    p["education"] = [{"institution": "○○大学", "degree": "修士",
                       "start": value, "end": "2023-03"}]
    fatal, _ = rr.date_problems(p)
    assert fatal == []


@pytest.mark.parametrize("value", REJECTED)
def test_unparseable_education_date_is_fatal_and_names_the_entry(jp_profile, value):
    p = dict(jp_profile)
    p["education"] = [{"institution": "○○大学", "degree": "修士",
                       "start": value, "end": "2023-03"}]
    fatal, _ = rr.date_problems(p)
    assert len(fatal) == 1
    assert "○○大学" in fatal[0]
    assert "start" in fatal[0]
    assert repr(value) in fatal[0] or "''" in fatal[0]


def test_a_current_role_needs_no_end_date(jp_profile):
    """現在に至る legitimately has no year. Flagging it would fire on every
    employed candidate."""
    fatal, _ = rr.date_problems(jp_profile)
    assert fatal == []


def test_an_undated_certification_warns_but_is_not_fatal(jp_profile):
    p = dict(jp_profile)
    p["certifications"] = ["基本情報技術者試験 合格"]
    fatal, warnings = rr.date_problems(p)
    assert fatal == []
    assert len(warnings) == 1 and "基本情報技術者試験" in warnings[0]


@pytest.mark.parametrize("cert,year,month", [
    ("基本情報技術者試験 合格 (2021)", "2021", ""),
    ("AWS Certified Cloud Practitioner (2024)", "2024", ""),
    ("普通自動車第一種運転免許 (2015-04)", "2015", "4"),
    ("2023年9月 応用情報技術者", "2023", "9"),
])
def test_a_certification_carrying_its_year_anywhere_is_read_and_silent(cert, year, month):
    """The most ordinary real entry puts the year in parentheses at the end.
    The anchored `_ym` returns ('', '') for it, which would blank the 年 cell
    AND fire a warning on a perfectly good line — a check that cries wolf on
    the common case is a check its reader stops seeing."""
    rows = rr.licenses_rows({"certifications": [cert]})
    assert (rows[0]["y"], rows[0]["m"]) == (year, month)
    fatal, warnings = rr.date_problems({"certifications": [cert]})
    assert fatal == [] and warnings == []


@pytest.mark.parametrize("cert", ["ISO 27001 Lead Auditor", "CCNA 200-301",
                                  "TOEIC 990"])
def test_a_digit_string_that_is_not_a_year_does_not_become_one(cert):
    """A false year printed into the 年 cell is worse than a blank one: it is
    a fabricated date on a legal-ish form, and it looks exactly like a real one."""
    rows = rr.licenses_rows({"certifications": [cert]})
    assert rows[0]["y"] == ""
    fatal, warnings = rr.date_problems({"certifications": [cert]})
    assert fatal == [] and len(warnings) == 1


def test_main_refuses_to_write_a_form_with_blank_year_cells(jp_profile, tmp_path, capsys):
    import yaml
    p = dict(jp_profile)
    p["education"] = [{"institution": "○○大学", "degree": "修士",
                       "start": "Sep 2023", "end": "2025-03"}]
    src = tmp_path / "profile.yaml"
    src.write_text(yaml.safe_dump(p, allow_unicode=True), encoding="utf-8")
    out = tmp_path / "rirekisho.md"
    assert rr.main([str(src), "--format", "md", "--out", str(out)]) == 1
    err = capsys.readouterr().err
    assert "WARNING" in err and "Sep 2023" in err
    assert not out.exists()


def test_allow_blank_dates_is_an_explicit_opt_in(jp_profile, tmp_path, capsys):
    import yaml
    p = dict(jp_profile)
    p["education"] = [{"institution": "○○大学", "degree": "修士",
                       "start": "Sep 2023", "end": "2025-03"}]
    src = tmp_path / "profile.yaml"
    src.write_text(yaml.safe_dump(p, allow_unicode=True), encoding="utf-8")
    out = tmp_path / "rirekisho.md"
    assert rr.main([str(src), "--format", "md", "--out", str(out),
                    "--allow-blank-dates"]) == 0
    assert out.exists()
    assert "WARNING" in capsys.readouterr().err     # still says so, just doesn't refuse


def test_a_fully_dated_profile_writes_and_is_silent(jp_profile, tmp_path, capsys):
    """The quiet case, pinned as hard as the firing one — including the
    certification, which is given its year here so that a silent run is a real
    claim about the whole form rather than about the tables we happened to fix."""
    import yaml
    p = dict(jp_profile)
    p["certifications"] = ["基本情報技術者試験 合格 (2021)"]
    src = tmp_path / "profile.yaml"
    src.write_text(yaml.safe_dump(p, allow_unicode=True), encoding="utf-8")
    out = tmp_path / "rirekisho.md"
    assert rr.main([str(src), "--format", "md", "--out", str(out)]) == 0
    assert "WARNING" not in capsys.readouterr().err
