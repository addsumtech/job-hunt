import pathlib
import sys
from docx import Document
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import render_cv


def test_reference_order_and_internship_group_survive_all_formats(tmp_path, capsys):
    profile = {
        "meta": {"name": "测试候选人", "language": "zh", "market": "cn",
                 "section_order": ["education", "experience", "internships", "projects", "skills"],
                 "headings": {"education": "教育背景", "experience": "工作经历", "internships": "实习经历（模板）", "projects": "个人项目与影响力", "skills": "技能"}},
        "education": [{"institution": "University A", "degree": "Master"}],
        "experience": [{"org": "Company A", "title": "Engineer", "bullets": ["First evidence", "Second evidence"]},
                       {"org": "Company B", "title": "Intern", "section": "internships", "bullets": ["Intern evidence"]}],
        "projects": [{"name": "Project A", "bullets": ["Project evidence"]}],
        "skills": {"technical": ["Python"]},
    }
    path = tmp_path / "cv.docx"
    render_cv.render_docx(profile, path)
    doc = Document(path)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["教育背景", "工作经历", "实习经历（模板）", "个人项目与影响力", "技能"]
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.index("First evidence") < text.index("Second evidence") < text.index("Company B")
    md = render_cv.render_markdown(profile)
    assert md.index("## 教育背景") < md.index("## 工作经历") < md.index("## 实习经历（模板）") < md.index("## 个人项目与影响力") < md.index("## 技能")
    assert "WARNING" not in capsys.readouterr().err
