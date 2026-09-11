"""Regressions from multilingual document delivery, with real render checks."""
import pathlib
import shutil
import sys

import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import check_pages
import deliver
import render_cv


@pytest.mark.parametrize('body', [
    '正文说明保留真实经历与岗位要求。',
    'Describe the actual project and the role requirements.',
    '実際の経験と応募要件を説明します。',
    '실제 경험과 채용 요건을 설명합니다.',
    'Describe la experiencia real y los requisitos del puesto.',
])
def test_portable_report_uses_reading_size_body_across_languages(tmp_path, body):
    import pymupdf
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text('# Report\n\n## Analysis\n\n' + body, encoding='utf-8')
    style = deliver._portable_report_style(md.read_text(encoding='utf-8'))
    assert style
    assert deliver._render_portable_report(md, pdf, style) == (True, '')
    with pymupdf.open(pdf) as document:
        spans = [span for page in document for block in page.get_text('dict')['blocks']
                 for line in block.get('lines', []) for span in line['spans']]
        matching = [span for span in spans if body[:5] in span['text']]
        assert matching and all(span['size'] == pytest.approx(12) for span in matching)
        assert max(span['size'] for span in spans) <= 18


def test_portable_report_flows_a_long_paragraph_without_losing_evidence(tmp_path):
    import pymupdf
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    tokens = [f'evidence{i:04d}' for i in range(800)]
    md.write_text('# Detailed report\n\n' + ' '.join(tokens), encoding='utf-8')
    style = deliver._portable_report_style(md.read_text(encoding='utf-8'))
    assert deliver._render_portable_report(md, pdf, style) == (True, '')
    with pymupdf.open(pdf) as document:
        assert len(document) > 1
        text = '\n'.join(page.get_text() for page in document)
        assert all(token in text for token in tokens)
        for page in document:
            assert all(page.rect.contains(block[:4]) for block in page.get_text('blocks'))


def test_standalone_official_link_is_not_printed_twice(tmp_path):
    import pymupdf
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text('# Roles\n\n[Beijing official posting](https://example.com/beijing)\n',
                  encoding='utf-8')
    style = deliver._portable_report_style(md.read_text(encoding='utf-8'))
    assert deliver._render_portable_report(md, pdf, style) == (True, '')
    with pymupdf.open(pdf) as document:
        assert document[0].get_text().count('Beijing official posting') == 1
        assert document[0].get_links()[0]['uri'] == 'https://example.com/beijing'


@pytest.mark.parametrize('title,anchor', [('Role analysis', 'role-analysis'),
                                         ('岗位分析', '岗位分析')])
def test_contents_resolves_forward_heading_links_after_pagination(tmp_path, title, anchor):
    import pymupdf
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text(f'# Report\n\n[First](#{anchor})\n\n[Second](#{anchor}-1)\n\n'
                  + ' '.join(f'evidence{i:04d}' for i in range(600)) +
                  f'\n\n## {title}\n\nFirst role.\n\n## {title}\n\nSecond role.', encoding='utf-8')
    style = deliver._portable_report_style(md.read_text(encoding='utf-8'))
    assert deliver._render_portable_report(md, pdf, style) == (True, '')
    with pymupdf.open(pdf) as document:
        links = document[0].get_links()
        assert len(links) == 2 and all(link['kind'] == pymupdf.LINK_GOTO for link in links)
        assert all(link['page'] > 0 for link in links)
        assert [link['page'] + 1 for link in links] == [entry[2] for entry in document.get_toc()]
        assert links[0]['to'].y < links[1]['to'].y


def test_unresolved_contents_link_is_not_delivered_as_a_broken_link(tmp_path):
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text('# Report\n\n[Missing section](#missing)', encoding='utf-8')
    style = deliver._portable_report_style(md.read_text(encoding='utf-8'))
    ok, reason = deliver._render_portable_report(md, pdf, style)
    assert not ok and 'unresolved report section link' in reason
    assert not pdf.exists()


def test_compact_comparison_table_keeps_columns_headers_and_links_across_pages(tmp_path):
    import pymupdf
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    urls = {f'https://example.com/jobs/{i}' for i in range(26)}
    rows = [f'| Role{i:02d} | Specific requirement {i}. [Official](https://example.com/jobs/{i}) |'
            for i in range(26)]
    md.write_text('# Roles\n\n| Role | Reason |\n|---|---|\n' + '\n'.join(rows),
                  encoding='utf-8')
    style = deliver._portable_report_style(md.read_text(encoding='utf-8'))
    assert deliver._render_portable_report(md, pdf, style) == (True, '')
    with pymupdf.open(pdf) as document:
        assert len(document) > 1
        actual = {link['uri'] for page in document for link in page.get_links()
                  if link['kind'] == pymupdf.LINK_URI}
        assert actual == urls
        for page in document:
            assert 'Reason' in page.get_text(), 'continued table lost its column header'
            reason = page.search_for('Reason')[0]
            role = page.search_for('Role')[0]
            assert reason.x0 > role.x0 + 100, 'comparison collapsed into a card title'
            fills = [drawing for drawing in page.get_drawings() if drawing.get('fill')]
            assert all(drawing['rect'].height < 50 for drawing in fills)


def test_wide_table_preserves_long_fields_and_city_links_without_card_frames(tmp_path):
    import pymupdf
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    tokens = [f'fact{i:04d}' for i in range(600)]
    urls = {'https://example.com/beijing', 'https://example.com/shanghai'}
    md.write_text('# Roles\n\n| Rank | Role | Evidence | Cities |\n|---|---|---|---|\n'
                  '| 1 | Agent PM | ' + ' '.join(tokens) +
                  ' | [Beijing](https://example.com/beijing), '
                  '[Shanghai](https://example.com/shanghai) |\n', encoding='utf-8')
    style = deliver._portable_report_style(md.read_text(encoding='utf-8'))
    assert deliver._render_portable_report(md, pdf, style) == (True, '')
    with pymupdf.open(pdf) as document:
        text = '\n'.join(page.get_text() for page in document)
        assert all(token in text for token in tokens)
        assert {link['uri'] for page in document for link in page.get_links()
                if link['kind'] == pymupdf.LINK_URI} == urls
        assert not any(drawing.get('fill') for page in document for drawing in page.get_drawings())


@pytest.mark.parametrize('layout', ['paragraph', 'compact_table', 'wide_table'])
def test_grouped_posting_links_each_occupy_their_own_line(tmp_path, layout):
    import pymupdf
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    urls = [f'https://example.com/jobs/{i}' for i in range(6)]
    links = ' '.join(f'[Beijing ID{i}]({url})' for i, url in enumerate(urls))
    if layout == 'compact_table':
        source = '| Role | Official entries |\n|---|---|\n| AI PM | ' + links + ' |'
    elif layout == 'wide_table':
        source = '| Rank | Role | Evidence | Links |\n|---|---|---|---|\n| 1 | AI PM | Summary | ' + links + ' |'
    else:
        source = links
    md.write_text('# Roles\n\n' + source, encoding='utf-8')
    style = deliver._portable_report_style(source)
    assert deliver._render_portable_report(md, pdf, style) == (True, '')
    with pymupdf.open(pdf) as document:
        positions = {}
        for page in document:
            for link in page.get_links():
                if link['kind'] == pymupdf.LINK_URI:
                    positions.setdefault(link['uri'], []).append((page.number, link['from']))
        assert set(positions) == set(urls)
        assert all(len(items) == 1 for items in positions.values())
        for first, second in zip(urls, urls[1:]):
            page_a, rect_a = positions[first][0]
            page_b, rect_b = positions[second][0]
            assert page_b > page_a or (page_b == page_a and rect_b.y0 >= rect_a.y1)


def test_selected_report_font_is_not_replaced_by_bundled_font(tmp_path, monkeypatch):
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text('中文报告', encoding='utf-8')
    seen = []
    monkeypatch.setattr(deliver, '_pandoc',
                        lambda m, p, f: (seen.append(f), p.write_bytes(b'PDF'), True)[-1])
    monkeypatch.setattr(deliver, '_verify_pdf', lambda *args: (True, ''))
    monkeypatch.setattr(deliver, '_render_portable_report',
                        lambda *args: pytest.fail('selected font was replaced'))
    assert deliver.render_pdf(md, pdf, 'SimSun') == (True, '')
    assert seen == ['SimSun']


def test_selected_report_font_failure_does_not_silently_substitute(tmp_path, monkeypatch):
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text('中文报告', encoding='utf-8')
    monkeypatch.setattr(deliver, '_pandoc', lambda *args: False)
    monkeypatch.setattr(deliver, '_render_portable_report',
                        lambda *args: pytest.fail('selected font was replaced'))
    ok, reason = deliver.render_pdf(md, pdf, 'Missing Font')
    assert not ok and 'selected report font' in reason and not pdf.exists()


def test_delivery_carries_profile_font_to_report(tmp_path, monkeypatch):
    ws, dest = tmp_path / 'workspace', tmp_path / 'delivery'
    ws.mkdir()
    (ws / 'report.md').write_text('中文报告', encoding='utf-8')
    (ws / 'tailored-profile.yaml').write_text('meta:\n  cjk_font: SimSun\n', encoding='utf-8')
    seen = []
    def render(md, pdf, font):
        seen.append(font)
        pdf.write_bytes(b'PDF')
        return True, ''
    monkeypatch.setattr(deliver, 'render_pdf', render)
    written, notes = deliver.deliver(ws, dest, 'candidate')
    assert not notes and len(written) == 2
    assert seen == ['SimSun']


@pytest.mark.parametrize('language,label', render_cv._PRESENT.items())
def test_current_dates_are_localized_without_changing_expected_graduation(language, label):
    profile = {'meta': {'name': 'Test Person', 'language': language},
               'experience': [{'title': 'Engineer', 'org': 'Example',
                               'start': '2023', 'end': 'present'}]}
    assert f'2023 – {label}' in render_cv.render_markdown(profile)
    assert render_cv._dates({'end': '2027 (expected)'}, language) == '2027 (expected)'
    assert render_cv._dates({'end': None}, language) == ''
    assert render_cv._dates_tex({'end': 'current'}, language) == label


@pytest.mark.skipif(not shutil.which('pandoc'), reason='requires pandoc parser')
def test_hidden_link_targets_are_not_visible_but_code_and_labels_are(tmp_path):
    md = tmp_path / 'report.md'
    md.write_text('[履歴](<https://example.org/隐藏中文路径>)\n\n`代码中文`\n', encoding='utf-8')
    text = deliver.visible_markdown(md)
    assert '履歴' in text and '代码中文' in text
    assert '隐藏中文路径' not in text


def test_wrong_cjk_glyphs_cannot_satisfy_coverage(tmp_path, monkeypatch):
    md, pdf = tmp_path / 'a.md', tmp_path / 'a.pdf'
    md.write_text('中文内容', encoding='utf-8')
    monkeypatch.setattr(deliver, 'visible_markdown', lambda p: '中文内容')
    monkeypatch.setattr(deliver, '_pandoc', lambda md, pdf, font: (pdf.write_bytes(b'PDF'), True)[1])
    monkeypatch.setattr(deliver, 'pdf_text', lambda p: '错误错误错误错误')
    monkeypatch.setattr(deliver, 'glyph_findings', lambda p: [])  # isolate text coverage
    ok, reason = deliver.render_pdf(md, pdf, 'fake')
    assert not ok and 'lost CJK' in reason and not pdf.exists()


def test_korean_probe_rejects_a_font_that_only_renders_chinese(tmp_path, monkeypatch):
    monkeypatch.setattr(deliver, '_pandoc', lambda *a: True)
    monkeypatch.setattr(deliver, 'pdf_text', lambda *a: '测试中文渲染')
    assert deliver.pick_cjk_font('한국어 문서') is None


def test_cjk_font_probe_stops_when_the_tex_engine_is_unavailable(monkeypatch):
    calls = []
    monkeypatch.setattr(deliver, '_pandoc', lambda *args: calls.append(args) or False)
    assert deliver.pick_cjk_font('中文报告') is None
    assert len(calls) == 1, "a broken engine must not be retried for every font"


def test_word_check_reports_missing_office_instead_of_passing(tmp_path, monkeypatch):
    doc = tmp_path / 'cv.docx'
    doc.write_bytes(b'not used')
    monkeypatch.setattr(check_pages.shutil, 'which', lambda name: None)
    assert check_pages.docx_findings(doc, {}, 2026)[0].startswith('DOCX_NOT_MEASURED:')


@pytest.mark.skipif(not shutil.which('soffice'), reason='requires LibreOffice')
def test_word_page_limit_measures_the_actual_word_file(tmp_path):
    from docx import Document
    profile = {'meta': {'name': 'Test Person', 'max_pages': 1},
               'experience': []}
    doc = Document()
    doc.add_paragraph('Test Person')
    doc.add_page_break()
    doc.add_paragraph('Education orphan on a second page')
    path = tmp_path / 'cv.docx'
    doc.save(path)
    findings = check_pages.docx_findings(path, profile, 2026)
    assert any('CV_TOO_LONG' in f and 'Word file cv.docx' in f for f in findings), findings


@pytest.mark.parametrize('text', ['한국어 문서 경력 기술 경험', '日本語の履歴書と経験',
                                 '中文报告与工作经历', '中文报告 日本語の履歴書 한국어 문서'])
def test_native_and_mixed_reports_really_round_trip_with_bundled_fonts(
        tmp_path, text, monkeypatch):
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text(text + '\n', encoding='utf-8')
    monkeypatch.setattr(deliver, '_pandoc',
                        lambda *args: pytest.fail('bundled CJK report took the TeX path'))
    ok, reason = deliver.render_pdf(md, pdf, None)
    assert ok, reason
    assert set(deliver._CJK.findall(text)) <= set(deliver._CJK.findall(deliver.pdf_text(pdf)))


@pytest.mark.skipif(not all(shutil.which(t) for t in ('pandoc', 'tectonic', 'pdftotext')),
                    reason='requires PDF toolchain')
def test_warning_symbol_has_a_readable_pdf_fallback_without_editing_source(tmp_path):
    md, pdf = tmp_path / 'notice.md', tmp_path / 'notice.pdf'
    source = '> ⚠️ Evidence count, not an outcome prediction.\n'
    md.write_text(source, encoding='utf-8')
    ok, reason = deliver.render_pdf(md, pdf, None)
    assert ok, reason
    text = deliver.pdf_text(pdf)
    assert '[!]' in text and 'Evidence count' in text
    assert md.read_text(encoding='utf-8') == source


@pytest.mark.parametrize('engine', ['tectonic', 'xelatex'])
def test_report_uses_available_supported_engine(tmp_path, monkeypatch, engine):
    import subprocess
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text('Report')
    monkeypatch.setattr(deliver.shutil, 'which', lambda name: '/bin/' + name if name in {'pandoc', engine} else None)
    seen = []
    def run(cmd, **kwargs):
        seen.append(cmd)
        pdf.write_bytes(b'%PDF test')
        return subprocess.CompletedProcess(cmd, 0)
    monkeypatch.setattr(deliver.subprocess, 'run', run)
    assert deliver._pandoc(md, pdf, None)
    assert '--pdf-engine=' + engine in seen[0]


def test_failed_report_compile_cannot_pass_on_partial_output(tmp_path, monkeypatch):
    import subprocess
    md, pdf = tmp_path / 'report.md', tmp_path / 'report.pdf'
    md.write_text('Report')
    monkeypatch.setattr(deliver.shutil, 'which', lambda name: '/bin/' + name)
    def fail(cmd, **kwargs):
        pdf.write_bytes(b'%PDF partial')
        return subprocess.CompletedProcess(cmd, 1)
    monkeypatch.setattr(deliver.subprocess, 'run', fail)
    # Portable report rendering deliberately avoids Pandoc. This regression
    # exercises the remaining Pandoc fallback, where a nonzero result must not
    # leave a partial output accepted as a finished PDF.
    monkeypatch.setattr(deliver, '_portable_report_style', lambda _source: None)
    ok, _ = deliver.render_pdf(md, pdf, None)
    assert not ok and not pdf.exists()


def test_cjk_does_not_select_luatex_for_xecjk(monkeypatch):
    import render_cv
    monkeypatch.setattr(render_cv, '_ENGINE_DIRS', ())
    monkeypatch.setattr(render_cv.shutil, 'which', lambda name: '/bin/lualatex' if name == 'lualatex' else None)
    assert render_cv.find_latex_engine(cjk=True) is None
    assert render_cv.find_latex_engine(cjk=False) == '/bin/lualatex'
