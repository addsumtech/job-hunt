#!/usr/bin/env python3
"""Render a Japanese 履歴書 (rirekisho) form from a canonical profile.

The rirekisho is the standardized Japanese CV form many traditional employers
expect — a different document from the Western CV that ``render_cv.py`` produces
(which maps to the companion 職務経歴書 / shokumu-keirekisho). It is form-based:

  • a personal-information block with a photo box,
  • a combined 学歴・職歴 (education + work history) chronological table,
  • a 免許・資格 (licenses / qualifications) table,
  • free-text 志望の動機 (motivation) and 本人希望記入欄 (requests) boxes.

Japan-specific fields live under ``jp:`` in the profile — all optional, and all
**personal data the candidate supplies** (never invented). The skill collects
them honestly:

  jp.name_furigana, jp.date_of_birth, jp.age, jp.gender, jp.address,
  jp.address_furigana, jp.photo_path, jp.date (the "as of" form date),
  jp.motivation, jp.personal_request

The 学歴・職歴 table is derived from the standard ``education`` and ``experience``
entries (merged, sorted, with 入学/卒業 and 入社/退社 rows). Nothing is fabricated.

Output: Markdown (preview) and ``.docx`` (the authentic, editable form — export
to PDF from Word/LibreOffice; the CJK form is not routed through LaTeX).

Usage:
  python render_rirekisho.py PROFILE.yaml --format md|docx --out OUTPATH
"""
import argparse
import pathlib
import re
import sys

sys.path.insert(0, str(pathlib.Path(__file__).parent))
import render_cv  # reuse load_profile + _is_current


# ── date helpers ──────────────────────────────────────────────────────────────

def _ym(s):
    """Extract (year, month) strings from a date like '2023-09' / '2021' /
    '2023年9月'. Returns ('', '') if no year is present."""
    m = re.match(r"\s*(\d{4})(?:[-/.年\s]+(\d{1,2}))?", str(s or ""))
    if not m:
        return "", ""
    month = str(int(m.group(2))) if m.group(2) else ""   # drop leading zero (9, not 09)
    return m.group(1), month


def _sortable(s):
    y, mo = _ym(s)
    return (int(y) if y else 9999, int(mo) if mo else 0)


def _checked_ym(value, what, problems):
    """(year, month) plus a recorded problem when there is no parseable year.

    _ym returns ('', '') for anything it cannot read — 'Sep 2023',
    'September 2023' and '03/2021' all do — and the row then renders with blank
    年 and 月 cells. That is a structurally invalid 履歴書 that renders, saves and
    passes pytest, and the rirekisho is deliberately routed away from all three
    judges, so no other reader exists.
    """
    y, mo = _ym(value)
    if not y and problems is not None:
        problems.append(f"{what}: could not read a year from {value!r} — accepted "
                        f"forms are '2023-09', '2023/09', '2021', '2023年9月'")
    return y, mo


# certifications[] is free-form prose, not a structured date field. The most
# ordinary real entry is "基本情報技術者試験 合格 (2021)", and the anchored _ym
# returns ('', '') for it — blanking the 年 cell while the year sits in plain
# sight, and firing a warning on a perfectly good line. So search instead of
# anchoring, and require a plausible year shape (19xx/20xx, not adjacent to
# another digit) so that "ISO 27001" and "CCNA 200-301" do not silently print
# a fabricated date into the form.
_CERT_YEAR = re.compile(r"(?<!\d)((?:19|20)\d{2})(?!\d)(?:\s*[-/.年]\s*(\d{1,2})(?!\d))?")


def _cert_ym(value):
    m = _CERT_YEAR.search(str(value or ""))
    if not m:
        return "", ""
    return m.group(1), (str(int(m.group(2))) if m.group(2) else "")


# ── content builders ──────────────────────────────────────────────────────────

def gakureki_shokureki_rows(profile, problems=None):
    """Build the combined 学歴・職歴 rows: a 学歴 header, enrolled/graduated rows
    per degree, a 職歴 header, joined/left rows per role, and a closing 以上."""
    rows = [{"y": "", "m": "", "text": "学歴", "align": "center"}]
    for ed in sorted(profile.get("education") or [], key=lambda e: _sortable(e.get("start"))):
        name = " ".join(x for x in [ed.get("institution", ""), ed.get("degree", "")] if x)
        y, mo = _checked_ym(ed.get("start"), f"学歴 {name} 入学 (start)", problems)
        rows.append({"y": y, "m": mo, "text": f"{name}　入学"})
        ey, emo = _checked_ym(ed.get("end"), f"学歴 {name} 卒業 (end)", problems) \
            if not render_cv._is_current(ed) else _ym(ed.get("end"))
        verb = "在学中" if render_cv._is_current(ed) else "卒業"
        rows.append({"y": ey, "m": emo, "text": f"{name}　{verb}"})

    rows.append({"y": "", "m": "", "text": "職歴", "align": "center"})
    exp = profile.get("experience") or []
    if not exp:
        rows.append({"y": "", "m": "", "text": "なし", "align": "center"})
    for ex in sorted(exp, key=lambda e: _sortable(e.get("start"))):
        org = ex.get("org", "")
        y, mo = _checked_ym(ex.get("start"), f"職歴 {org} 入社 (start)", problems)
        rows.append({"y": y, "m": mo, "text": f"{org}　入社"})
        if render_cv._is_current(ex):
            rows.append({"y": "", "m": "", "text": "現在に至る"})
        else:
            ey, emo = _checked_ym(ex.get("end"), f"職歴 {org} 退社 (end)", problems)
            rows.append({"y": ey, "m": emo, "text": f"{org}　退社"})

    rows.append({"y": "", "m": "", "text": "以上", "align": "right"})
    return rows


def licenses_rows(profile, problems=None):
    rows = []
    for cert in (profile.get("certifications") or []):
        y, mo = _cert_ym(cert)
        if not y and problems is not None:
            problems.append(f"免許・資格 {cert}: no year found — the 年 cell will be "
                            f"blank; add one, e.g. '基本情報技術者試験 合格 (2021)'")
        rows.append({"y": y, "m": mo, "text": cert})
    if not rows:
        rows.append({"y": "", "m": "", "text": "特になし"})
    return rows


def date_problems(profile):
    """(fatal, warnings). 学歴・職歴 rows must carry a year — the table IS the
    form, and a blank 年 cell makes it structurally invalid. A certification
    with no year at all is a common and tolerable omission, so it warns; a
    certification that carries its year in prose is read and stays silent."""
    fatal, warnings = [], []
    gakureki_shokureki_rows(profile, fatal)
    licenses_rows(profile, warnings)
    return fatal, warnings


# ── Markdown renderer (preview) ───────────────────────────────────────────────

def render_markdown(profile):
    jp = profile.get("jp") or {}
    meta = profile.get("meta") or {}
    c = profile.get("contact") or {}
    out = ["# 履歴書"]
    if jp.get("date"):
        out.append(f"{jp['date']}　現在")

    out += ["", "## 基本情報", ""]
    out.append(f"- ふりがな: {jp.get('name_furigana', '')}")
    out.append(f"- 氏名: {meta.get('name', '')}")
    dob, age = jp.get("date_of_birth", ""), jp.get("age", "")
    out.append(f"- 生年月日: {dob}" + (f"（満{age}歳）" if age else ""))
    if jp.get("gender"):
        out.append(f"- 性別: {jp['gender']}")
    if jp.get("address_furigana"):
        out.append(f"- ふりがな（現住所）: {jp['address_furigana']}")
    out.append(f"- 現住所: {jp.get('address', c.get('location', ''))}")
    out.append(f"- TEL: {c.get('phone', '')}")
    out.append(f"- E-mail: {c.get('email', '')}")

    out += ["", "## 学歴・職歴", "", "| 年 | 月 | 学歴・職歴 |", "|---|---|---|"]
    for r in gakureki_shokureki_rows(profile):
        out.append(f"| {r['y']} | {r['m']} | {r['text']} |")

    out += ["", "## 免許・資格", "", "| 年 | 月 | 免許・資格 |", "|---|---|---|"]
    for r in licenses_rows(profile):
        out.append(f"| {r['y']} | {r['m']} | {r['text']} |")

    if jp.get("motivation"):
        out += ["", "## 志望の動機", "", jp["motivation"].strip()]
    if jp.get("personal_request"):
        out += ["", "## 本人希望記入欄", "", jp["personal_request"].strip()]
    return "\n".join(out) + "\n"


# ── docx renderer (the authentic, editable form) ──────────────────────────────

def _ym_table(doc, title, rows):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc.add_heading(title, level=1)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "年", "月", title
    for r in rows:
        cells = t.add_row().cells
        if r.get("align") == "center":
            merged = cells[0].merge(cells[1]).merge(cells[2])
            merged.text = r["text"]
            merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            cells[0].text = str(r.get("y", ""))
            cells[1].text = str(r.get("m", ""))
            cells[2].text = r.get("text", "")
            if r.get("align") == "right":
                cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _box(doc, title, text):
    doc.add_heading(title, level=1)
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.cell(0, 0).text = text.strip()


def render_docx(profile, out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    jp = profile.get("jp") or {}
    meta = profile.get("meta") or {}
    c = profile.get("contact") or {}
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(15)
    section.left_margin = section.right_margin = Mm(18)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0
    for name, size in [("Title", 20), ("Heading 1", 12)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True
    doc.add_heading("履歴書", level=0)
    if jp.get("date"):
        p = doc.add_paragraph(f"{jp['date']}　現在")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Personal-info block: 4×2 grid, photo cell merged down the right column.
    info = doc.add_table(rows=4, cols=2)
    info.style = "Table Grid"
    info.cell(0, 0).text = f"ふりがな　{jp.get('name_furigana', '')}"
    info.cell(1, 0).text = f"氏名　{meta.get('name', '')}"
    dob, age = jp.get("date_of_birth", ""), jp.get("age", "")
    line = f"生年月日　{dob}" + (f"（満{age}歳）" if age else "")
    if jp.get("gender"):
        line += f"　　性別　{jp['gender']}"
    info.cell(2, 0).text = line
    addr = jp.get("address", c.get("location", ""))
    contact_text = ""
    if jp.get("address_furigana"):
        contact_text += f"ふりがな　{jp['address_furigana']}\n"
    contact_text += f"現住所　{addr}\nTEL　{c.get('phone', '')}　　E-mail　{c.get('email', '')}"
    info.cell(3, 0).text = contact_text

    photo = info.cell(0, 1).merge(info.cell(3, 1))
    photo_path = jp.get("photo_path")
    if photo_path and pathlib.Path(photo_path).exists():
        photo.paragraphs[0].add_run().add_picture(photo_path, width=Mm(30))
    else:
        photo.text = "写真貼付欄\n縦 36〜40mm\n横 24〜30mm"
    photo.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    _ym_table(doc, "学歴・職歴", gakureki_shokureki_rows(profile))
    _ym_table(doc, "免許・資格", licenses_rows(profile))
    if jp.get("motivation"):
        _box(doc, "志望の動機", jp["motivation"])
    if jp.get("personal_request"):
        _box(doc, "本人希望記入欄", jp["personal_request"])

    doc.save(str(out_path))


# ── CLI ───────────────────────────────────────────────────────────────────────

def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("profile")
    ap.add_argument("--format", choices=["md", "docx"], default="docx",
                    help="docx is the authentic form; export it to PDF from Word/LibreOffice")
    ap.add_argument("--out", required=True)
    ap.add_argument("--allow-blank-dates", action="store_true",
                    help="write the form anyway; the 年/月 cells will be blank")
    args = ap.parse_args(argv)

    profile = render_cv.load_profile(args.profile)
    fatal, warnings = date_problems(profile)
    for problem in fatal + warnings:
        print(f"WARNING: {problem}", file=sys.stderr)
    if fatal and not args.allow_blank_dates:
        print("Refusing to write a 履歴書 with blank 年/月 cells in 学歴・職歴 — the "
              "table is the form. Fix the dates, or pass --allow-blank-dates.",
              file=sys.stderr)
        return 1
    out = pathlib.Path(args.out)
    if args.format == "md":
        out.write_text(render_markdown(profile), encoding="utf-8")
    else:
        render_docx(profile, out)
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
